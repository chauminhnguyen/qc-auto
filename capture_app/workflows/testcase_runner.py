"""Execute generated workbook cases in isolated Playwright browser contexts."""

from __future__ import annotations

import argparse
import hashlib
import json
import re
import shutil
import unicodedata
import uuid
from datetime import datetime, timezone
from pathlib import Path
from urllib.parse import urlsplit

from openpyxl import load_workbook

from capture_app.core.ax_tree import Action, is_delete_like, is_submit_like
from capture_app.crawling import replayer
from capture_app.reporting.native_run_report import sanitize_text
from capture_app.workflows.case_metadata import load_metadata_sheet, row_hash
from capture_app.workflows.replay_url import normalize_replay_url

MUTATION_OBSERVER_JS = r"""
(() => {
  const key = "__testcaseObservationEvents";
  window[key] = window[key] || [];
  const selectors = [
    '[role="alert"]', '[role="status"]', '[aria-live]:not([aria-live="off"])',
    '[class*="toast" i]', '[class*="snackbar" i]', '[class*="notification" i]',
    '.field-validation-error', '.validation-summary-errors', '[aria-invalid="true"]'
  ].join(",");
  const add = node => {
    if (!node || node.nodeType !== 1) return;
    const candidates = node.matches?.(selectors) ? [node] : [];
    candidates.push(...(node.querySelectorAll?.(selectors) || []));
    for (const item of candidates) {
      const text = (item.innerText || item.textContent || "").trim();
      if (text) window[key].push({text, ts: Date.now(), role: item.getAttribute("role") || "",
        ariaLive: item.getAttribute("aria-live") || "", className: String(item.className || "")});
    }
  };
  new MutationObserver(records => records.forEach(record => {
    record.addedNodes.forEach(add);
    if (record.target?.nodeType === 1) add(record.target);
  })).observe(document, {subtree: true, childList: true, characterData: true});
})();
"""

JUDGE_SCHEMA = {
    "type": "object",
    "properties": {
        "verdict": {"type": "string", "enum": ["supported", "contradicted", "not_observed"]},
        "evidence_span": {"type": "string"},
    },
    "required": ["verdict", "evidence_span"],
    "additionalProperties": False,
}
JUDGE_PROMPT_VERSION = "claim-judge-v1"
DEFAULT_CALIBRATION_SET = Path(__file__).resolve().parents[2] / "tests" / "fixtures" / "semantic_calibration.jsonl"
JUDGE_SYSTEM = (
    "Compare one expected claim with only the supplied browser Observation. Return supported only when the "
    "observation directly supports it, contradicted only with explicit counter-evidence, otherwise "
    "not_observed. Return a verbatim evidence span from the observation, or an empty string. Do not infer "
    "from missing evidence. Prompt version: " + JUDGE_PROMPT_VERSION
)


def _configured_judge_model() -> str:
    from capture_app.workflows.testcase_gen import (
        ANTHROPIC_DEFAULT_MODEL, DEFAULT_MODEL, LLM_PROVIDER,
    )
    return ANTHROPIC_DEFAULT_MODEL if LLM_PROVIDER == "anthropic" else DEFAULT_MODEL

ERROR_MESSAGE_RE = re.compile(r"\b(error|failed|failure|invalid|rejected|declined|unable)\b|không (?:thành công|thể)|chưa (?:lưu|tạo|gửi)|lỗi", re.I)
NEGATED_SUCCESS_RE = re.compile(r"\b(?:not|never|no longer)\s+(?:saved|created|added|updated|successful|completed|submitted|confirmed|accepted|processed|sent|removed|deleted)\b", re.I)
SUCCESS_RE = re.compile(
    r"\b(saved|created|added|updated|success|successful|completed|submitted|confirmed|accepted|processed|sent|removed|deleted|approved|activated|enabled)\b"
    r"|đã\s+(?:lưu|tạo|thêm|cập nhật|gửi|xóa|thanh toán|xử lý|hoàn thành)"
    r"|(?:thành công|hoàn tất|hoàn thành)", re.I,
)
SENSITIVE_FIELD_RE = re.compile(r"password|token|secret|authorization|cookie|email|phone|card|cvv|address", re.I)
FIXTURE_RE = re.compile(r"\$\{([A-Za-z_][A-Za-z0-9_]*)\}|\{\{([A-Za-z_][A-Za-z0-9_]*)\}\}")


def _norm(value: object) -> str:
    return re.sub(r"\s+", " ", unicodedata.normalize("NFC", str(value or "")).casefold()).strip()


def _observation_search_text(value: object) -> str:
    if isinstance(value, dict):
        return " ".join(_observation_search_text(item) for item in value.values())
    if isinstance(value, (list, tuple)):
        return " ".join(_observation_search_text(item) for item in value)
    return str(value or "")


def _substitute_template(value: object, run_id: str, attempt_id: str,
                         fixtures: dict[str, str]) -> str:
    replacements = {"run_id": run_id, "attempt_id": attempt_id, **fixtures}
    return FIXTURE_RE.sub(lambda match: replacements.get(match.group(1) or match.group(2), match.group(0)),
                          str(value or ""))


def _origin(url: str) -> str:
    parts = urlsplit(str(url or ""))
    if parts.scheme not in {"http", "https"} or not parts.netloc:
        return ""
    return f"{parts.scheme.casefold()}://{parts.netloc.casefold()}"


def _redact(value: object, secrets: tuple[str, ...] = ()) -> str:
    return sanitize_text(str(value or ""), secrets).replace("\r", "")[:2000]


def _visible_row(sheet, row_number: int, headers: list[str]) -> dict:
    return {header: sheet.cell(row_number, index + 1).value
            for index, header in enumerate(headers)}


def _action_from_dict(item: dict, fixtures: dict[str, str], run_id: str,
                      attempt_id: str) -> Action:
    kind = str(item.get("kind") or "")
    name = str(item.get("name") or "")
    value = str(item.get("value") or "")
    if kind == "navigate":
        name = str(item.get("url") or name)
        kind = "goto"
    replacements = {"run_id": run_id, "attempt_id": attempt_id}
    missing = []

    def substitute(match):
        key = match.group(1) or match.group(2)
        if key in replacements:
            return replacements[key]
        if key in fixtures:
            return str(fixtures[key])
        missing.append(key)
        return "[MISSING_FIXTURE]"

    value = FIXTURE_RE.sub(substitute, value)
    name = FIXTURE_RE.sub(substitute, name)
    if kind == "goto":
        name = normalize_replay_url(name)
    if value.strip() in {"[TEST VALUE REQUIRED]", "[EXISTING VALUE REQUIRED]", "[REDACTED]"}:
        missing.append(str(item.get("name") or item.get("role") or "test data"))
    if missing:
        raise KeyError(sorted(set(missing))[0])
    return Action(kind, str(item.get("role") or ""), name, value,
                  int(item.get("dup_index") or 0))


def _fallback_claims(expected: str) -> list[dict]:
    fragments = [part.strip() for part in re.split(r"(?<=[.!?])\s+|\n+", expected) if part.strip()]
    claims = []
    for text in fragments[:4] or [expected.strip()]:
        quoted = re.search(r'["“](.+?)["”]', text)
        status = re.search(r"\b([2-5]xx)\b", text, re.I)
        url = re.search(r"https?://\S+|(?:^|\s)(/[A-Za-z0-9_./{}-]+)", text)
        field = re.search(r"(?:field|for)\s+[\"“]?([^\"”.,;]+)", text, re.I)
        claims.append({
            "text": text,
            "kind": "message" if quoted else "api_status" if status else "url" if url else "field_state" if field else "other",
            "literals": {
                "quoted_message_text": quoted.group(1) if quoted else None,
                "expected_status_class": status.group(1).lower() if status else None,
                "url_pattern": url.group(0).strip() if url else None,
                "field_name": field.group(1).strip() if field else None,
                "field_value": None,
            },
        })
    return claims


def _normalize_number_or_date(value: str) -> str:
    text = unicodedata.normalize("NFC", value).strip()
    date = re.fullmatch(r"(\d{1,2})/(\d{1,2})/(\d{4})", text)
    if date:
        day, month, year = date.groups()
        return f"{int(year):04d}-{int(month):02d}-{int(day):02d}"
    compact = text.replace("\u00a0", "").replace("\u202f", "").replace(" ", "")
    if re.fullmatch(r"[-+]?\d{1,3}(?:[.,\s]\d{3})+(?:[.,]\d+)?|[-+]?\d+(?:[.,]\d+)?", compact):
        if "," in compact and "." in compact:
            decimal = "," if compact.rfind(",") > compact.rfind(".") else "."
            thousands = "." if decimal == "," else ","
            compact = compact.replace(thousands, "").replace(decimal, ".")
        elif "," in compact:
            groups = compact.split(",")
            compact = "".join(groups) if len(groups) > 1 and all(len(g) == 3 for g in groups[1:]) else compact.replace(",", ".")
        elif compact.count(".") > 1:
            groups = compact.split(".")
            compact = "".join(groups) if all(len(g) == 3 for g in groups[1:]) else compact
        elif "." in compact and len(compact.rsplit(".", 1)[1]) == 3:
            compact = compact.replace(".", "")
        elif "," in compact and len(compact.rsplit(",", 1)[1]) == 3:
            compact = compact.replace(",", "")
        try:
            return f"{float(compact):.12g}"
        except ValueError:
            pass
    return _norm(text)


def _message_items(observation: dict) -> list[str]:
    source = observation.get("action_messages")
    if source is None:
        source = observation.get("messages", [])
    values = [str(item.get("text") if isinstance(item, dict) else item or "")
              for item in source if (item.get("text") if isinstance(item, dict) else item)]
    return values


def _primary_action_responses(observation: dict) -> list[dict]:
    responses = observation.get("own_request", {}).get("responses", [])
    mutating = [item for item in responses if str(item.get("method") or "").upper() in {"POST", "PUT", "PATCH", "DELETE"}]
    return mutating or responses


def _evaluate_structured(claim: dict, observation: dict) -> dict:
    kind = str(claim.get("kind") or "other")
    literals = claim.get("literals") or {}
    messages = _message_items(observation)
    normalized_messages = [(_norm(value), value) for value in messages]
    claim_text = str(claim.get("text") or "")
    contradictions = [message for message in messages
                      if ERROR_MESSAGE_RE.search(message) or NEGATED_SUCCESS_RE.search(message)]
    if contradictions and SUCCESS_RE.search(claim_text):
        return {"verdict": "contradicted", "evidence_span": contradictions[0], "basis": "structured", "outcome_type": "message"}
    if kind == "message":
        expected = str(literals.get("quoted_message_text") or "")
        expected_norm = _norm(expected)
        found = next((original for normalized, original in normalized_messages
                      if expected_norm and expected_norm in normalized), None)
        if found:
            return {"verdict": "supported", "evidence_span": found, "basis": "normalized", "outcome_type": "message"}
        return {"verdict": "not_observed", "evidence_span": "", "basis": "normalized", "outcome_type": ""}
    if kind == "api_status":
        expected_class = str(literals.get("expected_status_class") or "").lower()
        responses = _primary_action_responses(observation)
        response = next((item for item in responses if item.get("status") is not None), None)
        if response:
            actual_class = f"{int(response['status']) // 100}xx"
            evidence = f"{response.get('method', '')} {response.get('url', '')} -> {response['status']}"
            if not expected_class:
                return {"verdict": "not_observed", "evidence_span": "", "basis": "structured", "outcome_type": ""}
            verdict = "supported" if actual_class == expected_class else "contradicted"
            return {"verdict": verdict, "evidence_span": evidence, "basis": "structured", "outcome_type": "own_request"}
    if kind == "url":
        pattern = str(literals.get("url_pattern") or "")
        final_url = str(observation.get("url") or "")
        if pattern and final_url:
            if _norm(pattern) in _norm(final_url):
                return {"verdict": "supported", "evidence_span": final_url, "basis": "structured", "outcome_type": ""}
            return {"verdict": "contradicted", "evidence_span": final_url, "basis": "structured", "outcome_type": ""}
    if kind == "field_state":
        expected_name = _norm(literals.get("field_name"))
        expected_value = literals.get("field_value")
        for field in observation.get("field_states", []):
            if expected_name and expected_name in _norm(field.get("name")):
                actual = str(field.get("value") or field.get("state") or "")
                if expected_value is not None:
                    if _normalize_number_or_date(str(expected_value)) == _normalize_number_or_date(actual):
                        return {"verdict": "supported", "evidence_span": f"{field['name']}: {actual}", "basis": "structured", "outcome_type": "field_state"}
                    return {"verdict": "contradicted", "evidence_span": f"{field['name']}: {actual}", "basis": "structured", "outcome_type": "field_state"}
    normalized_claim = _norm(claim_text)
    for normalized, original in normalized_messages:
        if normalized_claim and normalized_claim in normalized:
            return {"verdict": "supported", "evidence_span": original, "basis": "normalized", "outcome_type": "message"}
    final_url = str(observation.get("url") or "")
    if final_url and literals.get("url_pattern") and _norm(literals["url_pattern"]) in _norm(final_url):
        return {"verdict": "supported", "evidence_span": final_url, "basis": "normalized", "outcome_type": ""}
    return {"verdict": "not_observed", "evidence_span": "", "basis": "normalized", "outcome_type": ""}


def _observation_snapshot(page, transient_start: int = 0) -> dict:
    snapshot = page.evaluate("""() => {
      const visible = el => !!(el.offsetWidth || el.offsetHeight || el.getClientRects().length);
      const text = el => (el.innerText || el.textContent || '').trim();
      const messages = Array.from(document.querySelectorAll('[role=alert],[role=status],[aria-live]:not([aria-live=off]),.toast,.snackbar,.notification'))
        .filter(visible).map(text).filter(Boolean);
      const observedAt = Date.now();
      const fields = Array.from(document.querySelectorAll('input,textarea,select,[role=checkbox],[role=switch]')).filter(visible).map(el => {
        const label = el.labels?.[0]?.innerText || el.getAttribute('aria-label') || el.name || el.id || '';
        return {name: label.trim(), value: (el.type === 'password' ? '[REDACTED]' : (el.value || '')),
          state: el.checked === undefined ? '' : (el.checked ? 'checked' : 'unchecked'), invalid: el.getAttribute('aria-invalid') === 'true'};
      });
      return {url: location.href, title: document.title || '', headings: Array.from(document.querySelectorAll('h1,h2,h3')).filter(visible).map(text).filter(Boolean),
        messages, message_events: messages.map(text => ({text, ts: observedAt, source: 'visible'})),
        field_states: fields, transient_messages: window.__testcaseObservationEvents || []};
    }""")
    return snapshot if isinstance(snapshot, dict) else {"url": str(getattr(page, "url", ""))}


def _redact_observation(observation: dict, secrets: tuple[str, ...]) -> dict:
    safe = dict(observation)
    for key in ("url", "title"):
        safe[key] = _redact(safe.get(key), secrets)
    for key in ("headings", "messages"):
        safe[key] = [_redact(item, secrets) for item in safe.get(key, [])]
    for key in ("message_events", "action_message_events"):
        safe[key] = [{**item, "text": _redact(item.get("text"), secrets)} for item in safe.get(key, [])]
    if "action_messages" in safe:
        safe["action_messages"] = [_redact(item, secrets) for item in safe["action_messages"]]
    safe["transient_messages"] = [
        {**item, "text": _redact(item.get("text"), secrets), "className": _redact(item.get("className"), secrets)}
        for item in safe.get("transient_messages", [])
    ]
    safe["field_states"] = [
        {**field, "name": _redact(field.get("name"), secrets),
         "value": "[REDACTED]" if SENSITIVE_FIELD_RE.search(str(field.get("name") or "")) else _redact(field.get("value"), secrets)}
        for field in safe.get("field_states", [])
    ]
    own = safe.get("own_request") or {"requests": [], "responses": [], "failures": []}
    safe["own_request"] = {
        key: [{**item, "url": _redact(item.get("url"), secrets),
               "failure": _redact(item.get("failure"), secrets) if item.get("failure") else item.get("failure")}
              for item in own.get(key, [])]
        for key in ("requests", "responses", "failures")
    }
    return safe


def _actual_result(observation: dict, error: str = "") -> str:
    lowered_error = str(error or "").casefold()
    if "browser startup failed" in lowered_error:
        return "The browser could not start, so this testcase could not be replayed."
    if "invalid url" in lowered_error:
        return (
            "The browser could not open the saved website address because it was invalid. "
            "It stayed on a blank page, so no messages or action response could be observed."
        )
    lines = []
    messages = _message_items(observation)
    if messages:
        lines.append("The page showed " + "; ".join(f'“{message}”' for message in messages) + ".")
    else:
        lines.append("No messages were captured.")
    title = str(observation.get("title") or "").strip()
    url = str(observation.get("url") or "").strip()
    if title and url:
        lines.append(f'The browser reached “{title}” at {url}.')
    elif url:
        lines.append(f"The browser ended at {url}.")
    elif title:
        lines.append(f'The browser reached “{title}”, but its URL was unavailable.')
    else:
        lines.append("The resulting page and URL were unavailable.")
    responses = observation.get("own_request", {}).get("responses", [])
    if responses:
        last = next((item for item in reversed(responses)
                     if item.get("method") in {"POST", "PUT", "PATCH", "DELETE"}), responses[-1])
        lines.append(
            f"The action received HTTP {last.get('status', 'unknown')} from "
            f"{last.get('method', 'request')} {last.get('url', 'an API endpoint')}."
        )
    else:
        lines.append("No same-origin API response was captured for the action.")
    headings = [_redact(value) for value in observation.get("headings", []) if str(value or "").strip()]
    if headings:
        lines.append("Headings: " + " | ".join(headings))
    fields = []
    for field in observation.get("field_states", []):
        name = _redact(field.get("name") or "field")
        value = "[REDACTED]" if SENSITIVE_FIELD_RE.search(name) else _redact(field.get("value") or field.get("state") or "")
        state = "invalid" if field.get("invalid") else "valid" if field.get("value") else ""
        fields.append(f"{name}={value}" + (f" ({state})" if state else ""))
    if fields:
        lines.append("Field states: " + " | ".join(fields))
    return "\n".join(lines)


def _literal_guard(claim: dict, verdict: dict, observation: dict, run_id: str, attempt_id: str) -> dict:
    if verdict.get("verdict") == "not_observed":
        return verdict
    normalized_observation = _norm(_observation_search_text(observation))
    evidence = unicodedata.normalize("NFC", str(verdict.get("evidence_span") or ""))
    if not evidence or _norm(evidence) not in normalized_observation:
        return {**verdict, "verdict": "not_observed", "evidence_span": "", "outcome_type": ""}
    if verdict.get("verdict") == "contradicted":
        return verdict
    literals = claim.get("literals") or {}
    replacements = {"run_id": run_id, "attempt_id": attempt_id}
    for name, value in literals.items():
        if value is None:
            continue
        literal = FIXTURE_RE.sub(lambda match: replacements.get(match.group(1) or match.group(2), match.group(0)), str(value))
        if name == "expected_status_class":
            actual_statuses = [int(item.get("status") or 0) for item in observation.get("own_request", {}).get("responses", [])]
            if not any(f"{status // 100}xx" == literal.casefold() for status in actual_statuses):
                return {**verdict, "verdict": "not_observed", "outcome_type": ""}
        elif name == "field_value":
            number_or_date = re.fullmatch(r"[-+]?\d[\d.,\s]*|\d{1,2}/\d{1,2}/\d{4}", literal)
            if number_or_date:
                actual_match = re.search(r"\d{1,2}/\d{1,2}/\d{4}|\d{1,3}(?:[.,\s]\d{3})+(?:[.,]\d+)?|\d+(?:[.,]\d+)?", evidence)
                if not actual_match or _normalize_number_or_date(literal) != _normalize_number_or_date(actual_match.group(0)):
                    return {**verdict, "verdict": "not_observed", "outcome_type": ""}
            elif _norm(literal) not in _norm(evidence):
                return {**verdict, "verdict": "not_observed", "outcome_type": ""}
        elif _norm(literal) not in _norm(evidence):
            return {**verdict, "verdict": "not_observed", "outcome_type": ""}
    return verdict


def _load_fixtures(path: Path | None) -> dict[str, str]:
    if path is None:
        return {}
    value = json.loads(path.read_text(encoding="utf-8"))
    if not isinstance(value, dict):
        raise ValueError("Fixture file must contain a JSON object.")
    return {str(key): str(item) for key, item in value.items()}


def _retryable_infrastructure_error(error: str) -> bool:
    lowered = error.casefold()
    return any(marker in lowered for marker in (
        "timeout", "execution context was destroyed", "navigation interrupted",
        "frame was detached", "target closed", "browser has been closed",
    ))


class TestcaseRunner:
    def __init__(self, workbook_path: str | Path, *, output_path: str | Path | None = None,
                 report_path: str | Path | None = None, fixtures: dict[str, str] | None = None,
                 auth_state: str | Path | None = None, allow_submit: bool = False,
                 allow_delete: bool = False, allowed_hosts: list[str] | None = None,
                 semantic: str = "advisory", only: str | None = None, resume: bool = False,
                 rerun_failed: bool = False, rerun_outcome_code: str | None = None,
                 allow_edited: bool = False, headed: bool = False,
                 browser_factory=None, judge=None):
        self.source_path = Path(workbook_path)
        self.output_path = Path(output_path) if output_path else self.source_path.with_name(self.source_path.stem + "-results" + self.source_path.suffix)
        if self.output_path.resolve() == self.source_path.resolve():
            raise ValueError("Results output must be a separate workbook copy.")
        self.report_path = Path(report_path) if report_path else self.output_path.with_suffix(".jsonl")
        self.fixtures = {str(k): str(v) for k, v in (fixtures or {}).items()}
        self.auth_state = Path(auth_state) if auth_state else None
        self.allow_submit = allow_submit
        self.allow_delete = allow_delete
        self.allowed_hosts = allowed_hosts or []
        self.semantic = semantic
        self.only = only
        self.resume = resume
        self.rerun_failed = rerun_failed
        self.rerun_outcome_code = rerun_outcome_code
        self.allow_edited = allow_edited
        self.headed = headed
        self.browser_factory = browser_factory
        self.judge = judge
        self.run_id = ""
        self._response_cache = None

    def _read_report(self) -> tuple[str, dict[str, dict]]:
        records = []
        if self.report_path.exists():
            with self.report_path.open("r", encoding="utf-8") as stream:
                for line in stream:
                    try:
                        item = json.loads(line)
                    except json.JSONDecodeError:
                        continue
                    if isinstance(item, dict):
                        records.append(item)
        header_index = next((index for index in range(len(records) - 1, -1, -1)
                             if records[index].get("type") == "header"), -1)
        header = records[header_index] if header_index >= 0 else None
        latest = {}
        run_id = str(header.get("run_id") or "") if header else ""
        for item in records[header_index + 1:]:
            if item.get("type") == "case" and item.get("case_id"):
                latest[str(item["case_id"])] = item
        return run_id, latest

    @staticmethod
    def _append_jsonl(path: Path, record: dict) -> None:
        path.parent.mkdir(parents=True, exist_ok=True)
        with path.open("a", encoding="utf-8", newline="\n") as stream:
            stream.write(json.dumps(record, ensure_ascii=False, separators=(",", ":")) + "\n")
            stream.flush()

    def _visible_for_hash(self, sheet, row_number: int, headers: list[str]) -> dict:
        return _visible_row(sheet, row_number, headers)

    def _origins(self, payload: dict) -> set[str]:
        found = {_origin(normalize_replay_url(payload.get("start_url")))}
        found.add(_origin(normalize_replay_url(payload.get("entry_url"))))
        for action in payload.get("route_actions", []):
            if isinstance(action, dict):
                raw_url = action.get("url") or (action.get("name") if action.get("kind") in {"goto", "navigate"} else "") or ""
                found.add(_origin(normalize_replay_url(raw_url)))
        checks = payload.get("observed_checks") or {}
        for key in ("before_url", "after_url"):
            found.add(_origin(str(checks.get(key) or "")))
        return {value for value in found if value}

    def _allowed_origins(self, payload: dict) -> set[str]:
        route_origins = self._origins(payload)
        start_origin = _origin(normalize_replay_url(payload.get("start_url")))
        entry_origin = _origin(normalize_replay_url(payload.get("entry_url")))
        origins = {origin for origin in (start_origin, entry_origin) if origin}
        explicit = set()
        for host in self.allowed_hosts:
            value = str(host).strip()
            if "://" in value:
                explicit.add(_origin(value))
            else:
                explicit.update(origin for origin in route_origins if urlsplit(origin).netloc.casefold() == value.casefold())
                if not any(urlsplit(origin).netloc.casefold() == value.casefold() for origin in route_origins):
                    explicit.add("https://" + value.casefold())
        return {value for value in origins | explicit if value}

    def _preflight(self, payload: dict, attempt_id: str | None = None) -> tuple[list[Action] | None, str]:
        actions = []
        safety_issues = []
        metadata_issues = []
        for item in payload.get("route_actions") or []:
            if not isinstance(item, dict):
                metadata_issues.append("invalid action metadata")
                continue
            kind = str(item.get("kind") or "")
            if kind == "navigate":
                kind = "goto"
            candidate = Action(kind, str(item.get("role") or ""),
                               str(item.get("url") or item.get("name") or ""),
                               str(item.get("value") or ""), int(item.get("dup_index") or 0))
            if is_delete_like(candidate) and not self.allow_delete:
                safety_issues.append("--allow-delete")
            if is_submit_like(candidate) and not self.allow_submit:
                safety_issues.append("--allow-submit")
        if safety_issues:
            return None, "HELD: route includes a gated action and requires " + " and ".join(sorted(set(safety_issues)) )
        try:
            if metadata_issues:
                return None, "REPLAY-ERROR: " + metadata_issues[0]
            start_url = str(payload.get("start_url") or "")
            entry_url = str(payload.get("entry_url") or "")
            missing_start = [match.group(1) or match.group(2) for match in FIXTURE_RE.finditer(start_url + " " + entry_url)
                             if (match.group(1) or match.group(2)) not in {"run_id", "attempt_id"}
                             and (match.group(1) or match.group(2)) not in self.fixtures]
            if missing_start:
                return None, f"NO-FIXTURE: missing fixture {missing_start[0]}"
            start_url = FIXTURE_RE.sub(lambda match: {"run_id": self.run_id}.get(
                match.group(1) or match.group(2), self.fixtures.get(match.group(1) or match.group(2), match.group(0))), start_url)
            entry_url = FIXTURE_RE.sub(lambda match: {"run_id": self.run_id}.get(
                match.group(1) or match.group(2), self.fixtures.get(match.group(1) or match.group(2), match.group(0))), entry_url)
            if "[REDACTED]" in start_url or "[REDACTED]" in entry_url:
                return None, "NO-FIXTURE: start URL contains a redacted value"
            start_url = normalize_replay_url(start_url)
            if not _origin(start_url):
                return None, "REPLAY-ERROR: start URL must include a valid host and scheme"
            entry_url = normalize_replay_url(entry_url) if entry_url else ""
            if entry_url and not _origin(entry_url):
                return None, "REPLAY-ERROR: testcase entry URL must include a valid host and scheme"
            if not entry_url:
                entry_url = self._infer_entry_url(payload, start_url)
            self._prepared_start_url = entry_url or start_url
            for item in payload.get("route_actions") or []:
                if not isinstance(item, dict):
                    return None, "REPLAY-ERROR: invalid action metadata"
                action = _action_from_dict(item, self.fixtures, self.run_id, attempt_id or str(uuid.uuid4()))
                actions.append(action)
        except KeyError as exc:
            return None, f"NO-FIXTURE: missing fixture {exc.args[0]}"
        if not payload.get("start_url"):
            return None, "REPLAY-ERROR: testcase metadata has no start URL"
        if not actions:
            return None, "REPLAY-ERROR: testcase metadata has no replay actions"
        return actions, ""

    @staticmethod
    def _infer_entry_url(payload: dict, start_url: str) -> str:
        """Use a captured source page when the stored route has no navigation path."""
        checks = payload.get("observed_checks") or {}
        source_url = normalize_replay_url(checks.get("before_url"))
        if not source_url or _origin(source_url) != _origin(start_url):
            return ""
        if urlsplit(source_url).path == urlsplit(start_url).path:
            return ""
        route_prefix = (payload.get("route_actions") or [])[:-1]
        if any(
            isinstance(action, dict)
            and str(action.get("kind") or "").casefold() in {"click", "navigate", "goto"}
            for action in route_prefix
        ):
            return ""
        return source_url

    def _new_context(self, browser):
        kwargs = {"service_workers": "block"}
        if self.auth_state:
            if not self.auth_state.is_file():
                raise FileNotFoundError("Storage-state file was not found.")
            kwargs["storage_state"] = str(self.auth_state)
        return browser.new_context(**kwargs)

    def _capture_case(self, payload: dict, actions: list[Action], attempt_id: str) -> tuple[dict, str]:
        if self.browser_factory:
            browser, close_browser, context = self.browser_factory(self.auth_state)
            close = close_browser
        else:
            from playwright.sync_api import sync_playwright
            playwright = sync_playwright().start()
            browser = playwright.chromium.launch(headless=not self.headed)
            context = self._new_context(browser)
            def close():
                try:
                    context.close()
                finally:
                    try:
                        browser.close()
                    finally:
                        playwright.stop()
        context.add_init_script(MUTATION_OBSERVER_JS)
        allowed = self._allowed_origins(payload)
        blocked_navigation = []

        def route_request(route):
            request = route.request
            try:
                if request.is_navigation_request() and _origin(request.url) not in allowed:
                    blocked_navigation.append(_redact(request.url, tuple(self.fixtures.values())))
                    route.abort()
                    return
                route.continue_()
            except Exception:
                try:
                    route.abort()
                except Exception:
                    pass

        context.route("**/*", route_request)
        page = context.new_page()
        self._action_origin = _origin(str(getattr(self, "_prepared_start_url", "") or payload.get("start_url") or page.url))
        page.on("request", lambda request: self._on_request(request, page))
        page.on("response", lambda response: self._on_response(response, page))
        page.on("requestfailed", lambda request: self._on_request_failed(request, page))
        self._active_request_capture = False
        self._action_requests = []
        self._action_responses = []
        self._action_failures = []
        visited = False
        error = ""
        try:
            replayer.goto_once(page, str(getattr(self, "_prepared_start_url", "") or payload["start_url"]))
            visited = True
            replayer.settle(page)
            replayer.wait_for_network_quiet(page)
            baseline_messages = _observation_snapshot(page).get("messages", [])
            action_started_at = 0
            for index, action in enumerate(actions):
                self._active_request_capture = index == len(actions) - 1
                if self._active_request_capture:
                    self._action_origin = _origin(str(getattr(page, "url", "")))
                    action_started_at = int(datetime.now(timezone.utc).timestamp() * 1000)
                    self._action_started_at = action_started_at
                replayer.perform_action(page, action)
                replayer.settle(page)
                replayer.wait_for_network_quiet(page)
            self._active_request_capture = False
            observation = _observation_snapshot(page)
            baseline = {_norm(item) for item in baseline_messages}
            transient_new = [item for item in observation.get("transient_messages", [])
                             if int(item.get("ts") or 0) >= action_started_at]
            visible_new = [item for item in observation.get("messages", []) if _norm(item) not in baseline]
            observation["action_messages"] = list(dict.fromkeys([*visible_new, *[
                item.get("text", "") for item in transient_new if item.get("text")
            ]]))
            event_time = int(datetime.now(timezone.utc).timestamp() * 1000)
            observation["action_message_events"] = [
                *[{"text": message, "ts": event_time, "source": "visible-after-action"} for message in visible_new],
                *transient_new,
            ]
            observation["transient_messages"] = transient_new
            observation["own_request"] = {
                "requests": self._action_requests,
                "responses": self._action_responses,
                "failures": self._action_failures,
            }
            observation.setdefault("action_messages", _message_items(observation))
        except Exception as exc:
            error = f"{type(exc).__name__}: {_redact(str(exc), tuple(self.fixtures.values()))[:250]}"
            observation = _observation_snapshot(page) if visited else {
                "url": _redact(str(getattr(page, "url", "")), tuple(self.fixtures.values())),
                "title": "", "headings": [], "messages": [], "field_states": [],
                "transient_messages": [],
            }
            observation["own_request"] = {
                "requests": self._action_requests,
                "responses": self._action_responses,
                "failures": self._action_failures,
            }
        finally:
            self._active_request_capture = False
            close()
        observation = _redact_observation(observation, tuple(self.fixtures.values()))
        if blocked_navigation and not error:
            error = "Blocked navigation outside --allowed-host."
        return observation, error

    def _on_request(self, request, page):
        try:
            if not getattr(self, "_active_request_capture", False):
                return
            if request.resource_type not in {"xhr", "fetch"} or _origin(request.url) != getattr(self, "_action_origin", _origin(page.url)):
                return
            self._action_requests.append({"method": str(request.method).upper(), "url": request.url})
        except Exception:
            return

    def _on_response(self, response, page):
        try:
            if not getattr(self, "_active_request_capture", False):
                return
            request = response.request
            if request.resource_type not in {"xhr", "fetch"} or _origin(response.url) != getattr(self, "_action_origin", _origin(page.url)):
                return
            self._action_responses.append({
                "method": str(request.method).upper(), "url": response.url,
                "status": int(response.status),
            })
        except Exception:
            return

    def _on_request_failed(self, request, page):
        try:
            if not getattr(self, "_active_request_capture", False):
                return
            if request.resource_type not in {"xhr", "fetch"} or _origin(request.url) != getattr(self, "_action_origin", _origin(page.url)):
                return
            self._action_failures.append({"method": str(request.method).upper(), "url": request.url,
                                          "failure": str(request.failure or "request failed")})
        except Exception:
            return

    def _semantic_judge(self, claim: dict, observation: dict) -> dict:
        safe_claim = dict(claim)
        safe_claim["text"] = _redact(safe_claim.get("text"), tuple(self.fixtures.values()))
        if isinstance(safe_claim.get("literals"), dict):
            safe_claim["literals"] = {
                key: (_redact(value, tuple(self.fixtures.values())) if value is not None else None)
                for key, value in safe_claim["literals"].items()
            }
        if self.judge:
            response = self.judge(safe_claim, observation)
        else:
            if self._response_cache is None:
                from capture_app.workflows.testcase_gen import _ResponseCache
                cache_path = self.report_path.with_name(self.report_path.stem + "-judge-cache.json")
                self._response_cache = _ResponseCache(cache_path, max_calls=100000)
            prompt = json.dumps({"claim": safe_claim, "observation": observation}, ensure_ascii=False, separators=(",", ":"))
            response = self._response_cache.chat_json(
                JUDGE_SYSTEM, prompt, model=None, max_tokens=350, response_schema=JUDGE_SCHEMA,
                temperature=0,
            )
        if not isinstance(response, dict) or response.get("verdict") not in {"supported", "contradicted", "not_observed"}:
            raw = json.dumps(response, ensure_ascii=False, default=str)
            raise ValueError("Judge response did not match the structured verdict schema; output=" +
                             _redact(raw, tuple(self.fixtures.values())))
        span = str(response.get("evidence_span") or "")
        judge_details = {
            "judge_input": {"claim": safe_claim, "observation": observation},
            "judge_output": {"verdict": response.get("verdict"),
                             "evidence_span": _redact(span, tuple(self.fixtures.values()))},
            "model": _configured_judge_model(), "prompt_version": JUDGE_PROMPT_VERSION,
        }
        if response["verdict"] != "not_observed":
            normalized_observation = _norm(_observation_search_text(observation))
            if not span or _norm(span) not in normalized_observation:
                return {"verdict": "not_observed", "evidence_span": "", "basis": "semantic", "outcome_type": "",
                        **judge_details}
        kind = str(claim.get("kind") or "other")
        outcome_type = ""
        if kind == "message" and span and any(_norm(span) in _norm(message) for message in _message_items(observation)):
            outcome_type = "message"
        elif kind == "field_state" and span and any(_norm(field.get("name")) in _norm(span)
                                                     for field in observation.get("field_states", [])):
            outcome_type = "field_state"
        elif kind == "api_status" and span and any(_norm(item.get("url")) in _norm(span)
                                                     for item in observation.get("own_request", {}).get("responses", [])):
            outcome_type = "own_request"
        return {"verdict": response["verdict"], "evidence_span": span,
                "basis": "semantic", "outcome_type": outcome_type,
                **judge_details}

    def _evaluate_case(self, expected: str, claims: list[dict], observation: dict,
                       attempt_id: str) -> tuple[str, list[dict], str]:
        own_responses = observation.get("own_request", {}).get("responses", [])
        primary_responses = _primary_action_responses(observation)
        if any(400 <= int(item.get("status") or 0) < 600 for item in primary_responses):
            failure = next(item for item in primary_responses if 400 <= int(item.get("status") or 0) < 600)
            veto = {"verdict": "contradicted", "evidence_span": f"{failure.get('method')} {failure.get('url')} -> {failure.get('status')}",
                    "basis": "structured", "outcome_type": "own_request"}
        else:
            veto = None
        results = []
        for claim in claims:
            templated_claim = dict(claim)
            templated_claim["text"] = _substitute_template(claim.get("text"), self.run_id, attempt_id, self.fixtures)
            templated_claim["literals"] = {
                name: (_substitute_template(value, self.run_id, attempt_id, self.fixtures) if value is not None else None)
                for name, value in (claim.get("literals") or {}).items()
            }
            result = _evaluate_structured(templated_claim, observation)
            result = _literal_guard(templated_claim, result, observation, self.run_id, attempt_id)
            if veto and SUCCESS_RE.search(str(templated_claim.get("text") or "")):
                result = dict(veto)
            if result["verdict"] == "not_observed" and self.semantic != "off":
                try:
                    semantic = self._semantic_judge(templated_claim, observation)
                    semantic = _literal_guard(templated_claim, semantic, observation, self.run_id, attempt_id)
                except Exception as exc:
                    raw_output = getattr(exc, "raw_output", None) or getattr(exc, "raw_response", None)
                    semantic = {"verdict": "not_observed", "evidence_span": "", "basis": "semantic",
                                "semantic_error": True, "error": type(exc).__name__,
                                "error_detail": _redact(str(exc), tuple(self.fixtures.values())),
                                "judge_input": {"claim": templated_claim, "observation": observation},
                                "judge_output": _redact(raw_output, tuple(self.fixtures.values())) if raw_output else _redact(str(exc), tuple(self.fixtures.values())),
                                "model": _configured_judge_model(),
                                "prompt_version": JUDGE_PROMPT_VERSION, "outcome_type": ""}
                result = semantic
            results.append({"claim": templated_claim, **result})
        determinate = [item for item in results if item["verdict"] != "not_observed"]
        deterministic_contradiction = any(
            item["verdict"] == "contradicted" and item.get("basis") != "semantic"
            for item in results
        )
        suggested = None
        if any(item.get("basis") == "semantic" for item in results):
            semantic_results = [item for item in results if item.get("basis") == "semantic"]
            all_verdicts = [item["verdict"] if item.get("basis") == "semantic" else item["verdict"] for item in results]
            suggested = "Failed" if "contradicted" in all_verdicts else "Passed" if all(v == "supported" for v in all_verdicts) else None
        if deterministic_contradiction:
            status = "Failed"
        elif self.semantic == "advisory":
            semantic_only = any(item.get("basis") == "semantic" for item in results)
            rollup = ["not_observed" if item.get("basis") == "semantic" else item["verdict"] for item in results]
            if all(value == "supported" for value in rollup) and any(item.get("outcome_type") for item in results if item.get("basis") != "semantic"):
                status = "Passed"
            elif any(value == "contradicted" for value in rollup):
                status = "Failed"
            else:
                status = "Unknown"
        elif results and all(item["verdict"] == "supported" for item in results) and any(item.get("outcome_type") for item in results):
            status = "Passed"
        elif any(item["verdict"] == "contradicted" for item in results):
            status = "Failed"
        else:
            status = "Unknown"
        if status == "Unknown" and not any(item["verdict"] != "not_observed" for item in results):
            code = "NO-ORACLE:"
        elif suggested and self.semantic == "advisory" and not deterministic_contradiction:
            code = "SEMANTIC-ONLY:"
        else:
            code = ""
        return status, results, f"{code} suggested={suggested}" if code and suggested else code

    def _run_cleanup(self, payload: dict) -> list[dict]:
        outcomes = []
        for cleanup in payload.get("cleanup") or []:
            cleanup_actions = cleanup.get("route_actions", [])
            run_tagged = bool(cleanup.get("tagged_with_run_id")) and any(
                "{{run_id}}" in str(action.get(key) or "") or "${run_id}" in str(action.get(key) or "")
                for action in cleanup_actions if isinstance(action, dict) for key in ("name", "url", "value")
            )
            if not self.allow_delete or not run_tagged:
                outcomes.append({"status": "HELD", "detail": "cleanup requires --allow-delete and current-run-id tagging"})
                continue
            # Cleanup uses a distinct context and the same allowed-origin policy.
            actions, issue = self._preflight({**payload, "route_actions": cleanup_actions}, str(uuid.uuid4()))
            if issue or not actions:
                outcomes.append({"status": "HELD", "detail": issue or "cleanup has no actions"})
                continue
            if not any(is_delete_like(action) for action in actions):
                outcomes.append({"status": "HELD", "detail": "cleanup route has no delete-like action"})
                continue
            try:
                observation, error = self._capture_case({**payload, "route_actions": cleanup["route_actions"]}, actions, str(uuid.uuid4()))
                outcomes.append({"status": "REPLAY-ERROR" if error else "completed", "detail": error,
                                 "observation": observation})
            except Exception as exc:
                outcomes.append({"status": "REPLAY-ERROR", "detail": type(exc).__name__})
        return outcomes

    def _case_prerequisite_issue(self, payload: dict, latest: dict[str, dict], case_id: str) -> str:
        for prerequisite in payload.get("prerequisites") or []:
            if isinstance(prerequisite, str):
                prerequisite = {"case_id": prerequisite}
            provider_id = str(prerequisite.get("case_id") or "")
            fixture_name = str(prerequisite.get("fixture") or "")
            if fixture_name and self.fixtures.get(fixture_name):
                continue
            provider = latest.get(provider_id)
            if provider and provider.get("case_id") != case_id and provider.get("action_request_succeeded"):
                continue
            return f"PREREQ-UNMET: provider {provider_id or '(unknown)'} has no successful action request in this run or fixture"
        return ""

    @staticmethod
    def _ordered_rows(by_row: dict[int, dict]) -> list[tuple[int, dict]]:
        """Keep workbook order where possible, moving in-workbook providers first."""
        pending = list(sorted(by_row.items()))
        case_rows = {str(meta.get("case_id")): row for row, meta in pending}
        ordered = []
        completed = set()
        while pending:
            ready_index = None
            for index, (_row, meta) in enumerate(pending):
                payload = meta.get("payload") or {}
                dependencies = []
                for prerequisite in payload.get("prerequisites") or []:
                    provider_id = str(prerequisite.get("case_id") or "") if isinstance(prerequisite, dict) else str(prerequisite)
                    if provider_id in case_rows and provider_id not in completed:
                        dependencies.append(provider_id)
                if not dependencies:
                    ready_index = index
                    break
            if ready_index is None:
                # Preserve deterministic order for dependency cycles. Each
                # dependent will be marked PREREQ-UNMET during execution.
                ordered.extend(pending)
                break
            item = pending.pop(ready_index)
            ordered.append(item)
            completed.add(str(item[1].get("case_id")))
        return ordered

    def run(self) -> dict:
        previous_run_id, latest = self._read_report()
        reuse_run = bool(self.resume or self.rerun_failed or self.rerun_outcome_code)
        if reuse_run and not previous_run_id:
            raise ValueError("Resume/rerun requires an existing JSONL report header.")
        self.run_id = previous_run_id if reuse_run else str(uuid.uuid4())
        if not reuse_run:
            latest = {}
        workbook_path = self.output_path if (self.resume or self.rerun_failed or self.rerun_outcome_code) and self.output_path.exists() else self.source_path
        if not self.output_path.exists():
            self.output_path.parent.mkdir(parents=True, exist_ok=True)
            shutil.copy2(self.source_path, self.output_path)
        workbook = load_workbook(workbook_path)
        if "Test Cases" not in workbook.sheetnames:
            raise ValueError("Workbook has no 'Test Cases' sheet.")
        sheet = workbook["Test Cases"]
        headers = {str(sheet.cell(5, col).value): col for col in range(1, sheet.max_column + 1)
                   if sheet.cell(5, col).value}
        required = {"Test Case Name", "Expected Result", "Actual Result", "Status", "Evidence/Note"}
        if not required.issubset(headers):
            raise ValueError("Workbook is missing testcase execution columns.")
        metadata = load_metadata_sheet(workbook)
        if not reuse_run:
            self._append_jsonl(self.report_path, {"type": "header", "run_id": self.run_id,
                                                  "created_at": datetime.now(timezone.utc).isoformat()})
        by_row = {item["sheet_row"]: item for item in metadata.values()}
        for row_number in range(6, sheet.max_row + 1):
            if row_number in by_row or not str(sheet.cell(row_number, headers["Test Case Name"]).value or "").strip():
                continue
            row_case_name = str(sheet.cell(row_number, headers["Test Case Name"]).value).strip()
            by_row[row_number] = {"case_id": row_case_name, "payload": None, "row_sha256": ""}
        # Classify every persisted route before Playwright launches. The per-row
        # preflight below repeats substitution with that run's attempt id.
        for meta in by_row.values():
            if meta.get("payload"):
                self._preflight(meta["payload"], "preflight")
                for cleanup in meta["payload"].get("cleanup") or []:
                    if isinstance(cleanup, dict):
                        self._preflight({**meta["payload"],
                                         "route_actions": cleanup.get("route_actions", [])}, "preflight-cleanup")
        outcomes: dict[str, int] = {}
        semantic_suggestions = 0
        cases_replayed = 0
        browser_start_error = None
        try:
            from playwright.sync_api import sync_playwright
            playwright_context = None
            browser = None
            try:
                for row_number, meta in self._ordered_rows(by_row):
                    case_id = meta["case_id"]
                    if self.only and self.only != case_id:
                        continue
                    existing = latest.get(case_id)
                    if existing:
                        if self.rerun_failed:
                            if existing.get("status") != "Failed" and existing.get("outcome_code") != "REPLAY-ERROR:":
                                continue
                        elif self.rerun_outcome_code:
                            if existing.get("outcome_code") != self.rerun_outcome_code:
                                continue
                        elif self.resume:
                            continue
                    if meta.get("payload") is None:
                        status, outcome_code = "Untested", "NO-FIXTURE:"
                        sheet.cell(row_number, headers["Actual Result"]).value = ""
                        sheet.cell(row_number, headers["Status"]).value = status
                        sheet.cell(row_number, headers["Evidence/Note"]).value = "NO-FIXTURE: no replay route/action metadata is available for this testcase."
                        record = {"type": "case", "run_id": self.run_id, "case_id": case_id,
                                  "status": status, "outcome_code": outcome_code, "actual_result": ""}
                        self._append_jsonl(self.report_path, record)
                        latest[case_id] = record
                        outcomes[outcome_code] = outcomes.get(outcome_code, 0) + 1
                        continue
                    row_values = self._visible_for_hash(sheet, row_number, list(headers.keys()))
                    actual_hash = row_hash(row_values, meta["payload"])
                    if actual_hash != meta["row_sha256"] and not self.allow_edited:
                        status, outcome_code = "Untested", "EDITED:"
                        sheet.cell(row_number, headers["Actual Result"]).value = ""
                        sheet.cell(row_number, headers["Status"]).value = status
                        sheet.cell(row_number, headers["Evidence/Note"]).value = outcome_code + " workbook row differs from generation metadata; not executed"
                        record = {"type": "case", "run_id": self.run_id, "case_id": case_id,
                                  "status": status, "outcome_code": outcome_code, "actual_result": ""}
                        self._append_jsonl(self.report_path, record)
                        latest[case_id] = record
                        outcomes[outcome_code] = outcomes.get(outcome_code, 0) + 1
                        continue
                    payload = meta["payload"]
                    if actual_hash != meta["row_sha256"] and self.allow_edited:
                        edited = True
                    else:
                        edited = False
                    prerequisite_issue = self._case_prerequisite_issue(payload, latest, case_id)
                    if prerequisite_issue:
                        status, outcome_code = "Untested", "PREREQ-UNMET:"
                        sheet.cell(row_number, headers["Actual Result"]).value = ""
                        sheet.cell(row_number, headers["Status"]).value = status
                        sheet.cell(row_number, headers["Evidence/Note"]).value = prerequisite_issue
                        record = {"type": "case", "run_id": self.run_id, "case_id": case_id,
                                  "status": status, "outcome_code": outcome_code, "actual_result": ""}
                        self._append_jsonl(self.report_path, record)
                        latest[case_id] = record
                        outcomes[outcome_code] = outcomes.get(outcome_code, 0) + 1
                        continue
                    attempt_id = str(uuid.uuid4())
                    actions, issue = self._preflight(payload, attempt_id)
                    if issue:
                        code = issue.split(":", 1)[0] + ":"
                        status = "Untested" if code in {"HELD:", "NO-FIXTURE:"} else "Unknown"
                        sheet.cell(row_number, headers["Actual Result"]).value = ""
                        sheet.cell(row_number, headers["Status"]).value = status
                        sheet.cell(row_number, headers["Evidence/Note"]).value = issue
                        record = {"type": "case", "run_id": self.run_id, "case_id": case_id,
                                  "status": status, "outcome_code": code, "actual_result": ""}
                        self._append_jsonl(self.report_path, record)
                        latest[case_id] = record
                        outcomes[code] = outcomes.get(code, 0) + 1
                        continue
                    try:
                        if self.browser_factory:
                            cases_replayed += 1
                            observation, error = self._capture_case(payload, actions, attempt_id)
                        else:
                            if browser_start_error:
                                raise RuntimeError(browser_start_error)
                            if playwright_context is None:
                                try:
                                    playwright_context = sync_playwright().start()
                                    browser = playwright_context.chromium.launch(headless=not self.headed)
                                except Exception as startup_exc:
                                    detail = _redact(startup_exc, tuple(self.fixtures.values()))
                                    browser_start_error = (
                                        f"Browser startup failed: {type(startup_exc).__name__}: "
                                        f"{detail or 'no error details available'}"
                                    )
                                    raise RuntimeError(browser_start_error) from startup_exc
                            cases_replayed += 1
                            def fresh_context(_auth_state):
                                context = self._new_context(browser)
                                return browser, context.close, context
                            self.browser_factory = fresh_context
                            observation, error = self._capture_case(payload, actions, attempt_id)
                            self.browser_factory = None
                        if error and not browser_start_error and _retryable_infrastructure_error(error):
                            retry_id = str(uuid.uuid4())
                            retry_actions, retry_issue = self._preflight(payload, retry_id)
                            if not retry_issue and retry_actions:
                                attempt_id = retry_id
                                try:
                                    if self.browser_factory:
                                        observation, error = self._capture_case(payload, retry_actions, retry_id)
                                    else:
                                        def retry_context(_auth_state):
                                            context = self._new_context(browser)
                                            return browser, context.close, context
                                        self.browser_factory = retry_context
                                        observation, error = self._capture_case(payload, retry_actions, retry_id)
                                        self.browser_factory = None
                                except Exception as retry_exc:
                                    error = f"{type(retry_exc).__name__}: browser retry failed"
                    except Exception as exc:
                        observation = {"url": "", "title": "", "headings": [], "messages": [],
                                       "transient_messages": [], "field_states": [],
                                       "own_request": {"requests": [], "responses": [], "failures": []}}
                        detail = _redact(exc, tuple(self.fixtures.values()))
                        error = f"{type(exc).__name__}: {detail or 'browser execution failed'}"
                    visible_actual = _actual_result(observation, error)
                    claims = payload.get("claims") or _fallback_claims(str(row_values.get("Expected Result") or ""))
                    if error:
                        status, claim_results, note = "Unknown", [], f"REPLAY-ERROR: {error}"
                    else:
                        status, claim_results, note = self._evaluate_case(
                            str(row_values.get("Expected Result") or ""), claims, observation, attempt_id,
                        )
                        if edited:
                            note = "EDITED: replay used stored metadata payload. " + note
                        elif status == "Unknown" and not note and any(
                                item.get("verdict") == "supported" for item in claim_results):
                            note = "Supported claims lacked qualifying outcome evidence."
                    if note.startswith("SEMANTIC-ONLY:"):
                        semantic_suggestions += 1
                    own_responses = _primary_action_responses(observation)
                    action_request_class = ""
                    if own_responses:
                        priority_responses = [item for item in own_responses if item.get("method") in {"POST", "PUT", "PATCH", "DELETE"}] or own_responses
                        code = int(priority_responses[0].get("status") or 0) // 100
                        action_request_class = f"{code}xx" if code else ""
                    evidence = [note] if note else []
                    for result in claim_results:
                        evidence.append(f"{result['verdict']} basis={result.get('basis')} evidence={result.get('evidence_span')}")
                    safe_claim_results = []
                    for result in claim_results:
                        safe_result = dict(result)
                        if isinstance(safe_result.get("claim"), dict):
                            safe_result["claim"] = {
                                **safe_result["claim"],
                                "text": _redact(safe_result["claim"].get("text"), tuple(self.fixtures.values())),
                                "literals": {key: (_redact(value, tuple(self.fixtures.values())) if value is not None else None)
                                             for key, value in (safe_result["claim"].get("literals") or {}).items()},
                            }
                        if isinstance(safe_result.get("judge_input"), dict) and isinstance(
                                safe_result["judge_input"].get("claim"), dict):
                            judge_claim = safe_result["judge_input"]["claim"]
                            safe_result["judge_input"]["claim"] = {
                                **judge_claim,
                                "text": _redact(judge_claim.get("text"), tuple(self.fixtures.values())),
                                "literals": {key: (_redact(value, tuple(self.fixtures.values())) if value is not None else None)
                                             for key, value in (judge_claim.get("literals") or {}).items()},
                            }
                        safe_claim_results.append(safe_result)
                    sheet.cell(row_number, headers["Actual Result"]).value = visible_actual
                    sheet.cell(row_number, headers["Status"]).value = status
                    sheet.cell(row_number, headers["Evidence/Note"]).value = "\n".join(evidence)
                    cleanup_records = self._run_cleanup(payload)
                    record = {
                        "type": "case", "run_id": self.run_id, "case_id": case_id,
                        "status": status, "outcome_code": ("REPLAY-ERROR:" if error else note.split(" ", 1)[0] if note.startswith(("NO-ORACLE:", "SEMANTIC-ONLY:")) else ""),
                        "action_request_class": action_request_class,
                        "action_request_succeeded": any(200 <= int(item.get("status") or 0) < 300 for item in own_responses),
                        "actual_result": visible_actual, "observation": observation,
                        "claim_results": safe_claim_results, "cleanup": cleanup_records,
                        "edited_metadata_used": edited,
                    }
                    self._append_jsonl(self.report_path, record)
                    latest[case_id] = record
                    code = record["outcome_code"] or status
                    outcomes[code] = outcomes.get(code, 0) + 1
            finally:
                if browser:
                    browser.close()
                if playwright_context:
                    playwright_context.stop()
        finally:
            workbook.save(self.output_path)
            workbook.close()
        summary = {"run_id": self.run_id, "output_path": str(self.output_path),
                   "report_path": str(self.report_path), "outcomes": outcomes,
                   "semantic_suggestions": semantic_suggestions,
                   "cases_replayed": cases_replayed,
                   "browser_start_error": browser_start_error if not self.browser_factory else None}
        print("[testcase_runner] " + json.dumps(summary, ensure_ascii=False))
        return summary


def _run_calibration(path: Path, judge=None) -> dict:
    rows = [json.loads(line) for line in path.read_text(encoding="utf-8").splitlines() if line.strip()]
    results = []
    runner = TestcaseRunner(path, semantic="on", judge=judge)
    for row in rows:
        claim = row["claim"]
        observation = row["observation"]
        verdict = runner._semantic_judge(claim, observation)
        expected = str(row.get("label") or "").casefold()
        results.append({"label": expected, "verdict": verdict["verdict"]})
    false_passes = sum(item["verdict"] == "supported" and item["label"] != "passed" for item in results)
    negatives = sum(item["label"] != "passed" for item in results)
    report = {"rows": len(results), "false_passes": false_passes,
              "negative_rows": negatives, "false_pass_rate": false_passes / negatives if negatives else 0.0,
              "model": _configured_judge_model(), "prompt_version": JUDGE_PROMPT_VERSION}
    print(json.dumps(report, ensure_ascii=False))
    return report


def main(argv: list[str] | None = None) -> int:
    parser = argparse.ArgumentParser(description="Run generated testcases in isolated Playwright contexts.")
    parser.add_argument("--workbook", type=Path)
    parser.add_argument("--out", type=Path)
    parser.add_argument("--report", type=Path)
    parser.add_argument("--fixtures", type=Path)
    parser.add_argument("--auth-state", type=Path)
    parser.add_argument("--allow-submit", action="store_true")
    parser.add_argument("--allow-delete", action="store_true")
    parser.add_argument("--allowed-host", action="append", default=[])
    parser.add_argument("--semantic", choices=("off", "advisory", "on"), default="advisory")
    parser.add_argument("--only", metavar="CASE_ID")
    parser.add_argument("--resume", action="store_true")
    parser.add_argument("--rerun-failed", action="store_true")
    parser.add_argument("--rerun-outcome-code")
    parser.add_argument("--allow-edited", action="store_true")
    parser.add_argument("--headed", action="store_true")
    parser.add_argument("--calibrate", nargs="?", type=Path, const=DEFAULT_CALIBRATION_SET,
                        metavar="LABELED_JSONL", help="judge the labeled calibration set; rerun after model or prompt changes")
    args = parser.parse_args(argv)
    if args.calibrate:
        _run_calibration(args.calibrate)
        return 0
    if not args.workbook:
        parser.error("--workbook is required unless --calibrate is used")
    try:
        fixtures = _load_fixtures(args.fixtures)
        runner = TestcaseRunner(
            args.workbook, output_path=args.out, report_path=args.report, fixtures=fixtures,
            auth_state=args.auth_state, allow_submit=args.allow_submit, allow_delete=args.allow_delete,
            allowed_hosts=args.allowed_host, semantic=args.semantic, only=args.only,
            resume=args.resume, rerun_failed=args.rerun_failed,
            rerun_outcome_code=args.rerun_outcome_code, allow_edited=args.allow_edited,
            headed=args.headed,
        )
        runner.run()
    except (OSError, ValueError, json.JSONDecodeError) as exc:
        parser.error(str(exc))
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
