import json
import sqlite3
import sys
import tempfile
import unittest
from pathlib import Path
from types import SimpleNamespace
from unittest.mock import patch

from openpyxl import load_workbook

from capture_app.workflows import testcase_gen


def make_vault(root: Path, *, sessions=("s1",), transition_rows=150, include_transitions=True):
    conn = sqlite3.connect(":memory:")
    conn.execute("CREATE TABLE sessions(session_id TEXT PRIMARY KEY, start_url TEXT, started_at TEXT)")
    conn.execute("""CREATE TABLE steps (
        session_id TEXT, seq INTEGER, step_type TEXT, selector TEXT, label TEXT, value TEXT,
        input_type TEXT, page_url TEXT, tag TEXT, elem_type TEXT, role TEXT, checked INTEGER,
        is_native_control INTEGER, label_for TEXT, select_kind TEXT, selected_options_json TEXT, ts REAL,
        action_id TEXT
    )""")
    conn.execute("""CREATE TABLE session_api_events (
        session_id TEXT, seq INTEGER, method TEXT, status INTEGER, request_url TEXT,
        action_id TEXT, correlation_status TEXT, entry_point TEXT
    )""")
    for session_id in sessions:
        conn.execute("INSERT INTO sessions VALUES (?, ?, ?)", (session_id, "https://example.test/home", session_id))
        conn.execute(
            "INSERT INTO steps VALUES (?, 1, 'tap', '#save', 'Save', '', '', "
            "'https://example.test/home', 'button', 'button', 'button', NULL, 0, '', '', NULL, 1, 'a1')",
            (session_id,),
        )
        conn.execute(
            "INSERT INTO session_api_events VALUES (?, 1, 'POST', 201, '/api/save', 'a1', 'matched', '')",
            (session_id,),
        )

    if include_transitions:
        conn.execute("""CREATE TABLE ui_transitions (
            crawl_id TEXT, seq INTEGER, from_fp TEXT, to_fp TEXT, to_url TEXT,
            action_kind TEXT, action_role TEXT, action_name TEXT, action_value TEXT,
            action_dup_index INTEGER, path_json TEXT, triggered_apis_json TEXT,
            ok INTEGER, detail TEXT, start_url TEXT, session_id TEXT, from_url TEXT, ts REAL
        )""")
        for session_id in sessions:
            for seq in range(1, transition_rows + 1):
                conn.execute(
                    "INSERT INTO ui_transitions VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?)",
                    (
                        "run-latest", seq, "home", f"dest-{seq}", f"https://example.test/page/{seq}",
                        "click", "button", f"Action {seq}", "", 0, "[]",
                        json.dumps([{"method": "POST", "url": "/api/action", "status": 200 if seq % 2 else 500}]),
                        int(seq % 2 == 1), "observed", "https://example.test/home", session_id,
                        "https://example.test/home", float(seq),
                    ),
                )
    conn.commit()
    return SimpleNamespace(conn=conn, vault_path=root)


class LiveRunTestcaseGenerationTests(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory(prefix="testcase-live-run-")
        self.addCleanup(self.temp.cleanup)
        self.root = Path(self.temp.name)

    def test_generates_happy_case_and_all_150_ordered_transitions(self):
        vault = make_vault(self.root)
        rows = testcase_gen.generate_from_live_run(vault, "s1")

        self.assertEqual(len(rows), 151)
        self.assertEqual(rows[0]["Category"], "happy-path")
        self.assertIn("POST /api/save (HTTP 201)", rows[0]["Expected Result"])
        self.assertIn("Live run run-latest transition 1:", rows[1]["Test Case Name"])
        self.assertIn("HTTP 500", rows[2]["Expected Result"])
        self.assertIn("Live run run-latest transition 150:", rows[-1]["Test Case Name"])
        self.assertEqual(sum("recorded live-run API outcome" in row["Expected Result"] for row in rows[1:]), 150)

    def test_structured_loader_isolates_session_and_honors_requested_run(self):
        vault = make_vault(self.root, sessions=("s1", "s2"), transition_rows=2)
        vault.conn.execute("UPDATE ui_transitions SET crawl_id='run-other' WHERE session_id='s2'")
        vault.conn.execute("UPDATE ui_transitions SET crawl_id='run-old', ts=0 WHERE session_id='s1' AND seq=1")
        vault.conn.commit()

        selected, rows = testcase_gen._load_persisted_live_run(vault, "s1", "run-latest", include_legacy=False)
        self.assertEqual(selected, "run-latest")
        self.assertEqual([row["seq"] for row in rows], [2])
        with self.assertRaisesRegex(ValueError, "No persisted transition data"):
            testcase_gen._load_persisted_live_run(vault, "s1", "run-other", include_legacy=False)

    def test_legacy_chat_evidence_fallback_loads_transition(self):
        vault = make_vault(self.root, transition_rows=0, include_transitions=False)
        chat_dir = self.root / ".chat_index"
        chat_dir.mkdir()
        conn = sqlite3.connect(chat_dir / "index.sqlite")
        conn.execute("CREATE TABLE chunks(source_path TEXT, body TEXT, note_type TEXT)")
        conn.execute("""INSERT INTO chunks VALUES (?, ?, ?)""", (
            "live_runner/old-run/transition-1",
            "Session: s1\nStatus: succeeded\nFrom UI state: home\nTo UI state: destination\n"
            "Destination URL: https://example.test/next\nAction: Click button Continue\n"
            "Outcome detail: complete\nAction path:\n1. Click button Continue\nTriggered APIs:\n"
            "- POST /api/continue; status=200",
            "live-run-transition",
        ))
        conn.commit()
        conn.close()

        rows = testcase_gen.generate_from_live_run(vault, "s1")
        self.assertEqual(len(rows), 2)
        self.assertIn("old-run transition 1", rows[1]["Test Case Name"])
        self.assertIn("POST /api/continue (HTTP 200)", rows[1]["Expected Result"])

    def test_malformed_records_are_rejected_before_workbook_write(self):
        vault = make_vault(self.root, transition_rows=1)
        vault.conn.execute("UPDATE ui_transitions SET path_json='{' WHERE seq=1")
        vault.conn.commit()

        with self.assertRaisesRegex(ValueError, "Malformed live-run transition data"):
            testcase_gen.generate_from_live_run(vault, "s1")

    def test_missing_session_does_not_replace_existing_workbook(self):
        vault = make_vault(self.root, sessions=(), transition_rows=0)
        output = self.root / "testcases.xlsx"
        output.write_bytes(b"existing workbook stays intact")
        with patch.object(testcase_gen, "VaultWriter", return_value=vault), patch.object(
            sys, "argv", ["testcase_gen", "--vault", str(self.root), "--out", str(output)],
        ):
            with self.assertRaises(SystemExit):
                testcase_gen.main()
        self.assertEqual(output.read_bytes(), b"existing workbook stays intact")

    def test_workbook_keeps_11_columns_and_untested_status(self):
        vault = make_vault(self.root, transition_rows=1)
        rows = testcase_gen.generate_from_live_run(vault, "s1")
        output = testcase_gen.write_excel(rows, self.root / "cases.xlsx")
        workbook = load_workbook(output, data_only=False)
        sheet = workbook["Test Cases"]

        self.assertEqual(sheet.max_column, 11)
        self.assertEqual(sheet.max_row, 7)
        self.assertEqual(sheet.cell(5, 1).value, "Module")
        self.assertEqual(sheet.cell(5, 11).value, "Evidence/Note")
        self.assertEqual(sheet.cell(6, 10).value, "Untested")
        self.assertEqual(sheet.cell(7, 10).value, "Untested")
        workbook.close()


if __name__ == "__main__":
    unittest.main()
