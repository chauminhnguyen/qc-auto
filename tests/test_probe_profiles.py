import tempfile
import unittest
from pathlib import Path
from types import SimpleNamespace
from unittest.mock import patch

from capture_app.storage.vault import MAX_BASELINE_RESPONSE_CHARS, VaultWriter
from capture_app.workflows import probe_profiles
from capture_app.workflows import testcase_gen
from capture_app.workflows import enrich


class ProbeProfileTests(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory(prefix="probe-profile-")
        self.addCleanup(self.temp.cleanup)
        self.vault = VaultWriter(Path(self.temp.name))
        self.addCleanup(self.vault.conn.close)
        self.api = {
            "id": 1, "ref_stem": "POST_items", "method": "POST",
            "full_url": "https://example.test/items", "headers_json": '{"Content-Type":"application/json"}',
            "payload_json": '{"name":"known"}',
            "captured_response_status": 200, "captured_response_body": "baseline",
        }

    def test_baseline_failure_stops_before_field_probes(self):
        with patch.object(probe_profiles, "_send", return_value=(401, '{"message":"expired"}')) as send, \
             patch.object(probe_profiles.time, "sleep"):
            result = probe_profiles.probe_api(
                self.vault, "s::r", self.api, "r:1", [{"name": "name", "baseline": "known"}],
                confirm_state_change=True,
            )
        self.assertEqual(result["status"], "baseline_failed")
        self.assertEqual(send.call_count, 1)
        rows = self.vault.get_field_probe_log("s::r")
        self.assertEqual([row["probe_kind"] for row in rows], ["baseline"])

    def test_card_and_password_fields_skip_api_before_any_request(self):
        for field_name in ("CardNumber", "Password"):
            with self.subTest(field_name=field_name), \
                 patch.object(probe_profiles, "_send") as send:
                result = probe_profiles.probe_api(
                    self.vault, "s::r", self.api, "r:1",
                    [{"name": field_name, "baseline": "synthetic"}],
                    confirm_state_change=True,
                )
            self.assertEqual(result["status"], "sensitive_fields_skipped")
            send.assert_not_called()
            self.assertEqual(self.vault.get_field_probe_log("s::r"), [])

    def test_each_probe_changes_only_the_mapped_field(self):
        seen = []

        def fake_send(api, values=None, timeout=30.0):
            seen.append(values)
            return (201, '{"ok":true}')

        with patch.object(probe_profiles, "_send", side_effect=fake_send), \
             patch.object(probe_profiles.time, "sleep"):
            probe_profiles.probe_api(
                self.vault, "s::r", self.api, "r:1", [{"name": "name", "baseline": "known"}],
                confirm_state_change=True,
            )
        self.assertTrue(seen)
        self.assertTrue(all(set(payload) == {"name"} for payload in seen[1:]))

    def test_value_only_agreement_does_not_match_a_field(self):
        class ParamsVault:
            def get_params_for_api(self, _ref, role=None):
                return [{"name": "email", "examples": ["same-value"]}]

        transition = {"path_json": '[{"kind":"fill","role":"textbox",'
                                   '"name":"textbox","value":"same-value"}]'}
        result = probe_profiles.match_transition_inputs(
            ParamsVault(), transition, {"ref_stem": "POST_items", "payload_json": "{}"},
        )
        self.assertEqual(result[0]["status"], "unknown")
        self.assertIn("value-only matches are not accepted", result[0]["reason"])

    def test_duplicate_name_params_are_reported_ambiguous(self):
        class ParamsVault:
            def get_params_for_api(self, _ref, role=None):
                return [
                    {"name": "email", "examples": ["one@example.test"]},
                    {"name": "e-mail", "examples": ["one@example.test"]},
                ]

        transition = {"path_json": '[{"kind":"fill","role":"textbox",'
                                   '"name":"Email","value":"one@example.test"}]'}
        result = probe_profiles.match_transition_inputs(
            ParamsVault(), transition, {"ref_stem": "POST_items", "payload_json": "{}"},
        )
        self.assertEqual(result[0]["status"], "ambiguous")
        self.assertEqual(set(result[0]["candidate_params"]), {"email", "e-mail"})

    def test_value_can_support_but_not_create_a_name_match(self):
        class ParamsVault:
            def get_params_for_api(self, _ref, role=None):
                return [
                    {"name": "email", "examples": ["same-value"]},
                    {"name": "e-mail", "examples": ["other-value"]},
                ]

        transition = {"path_json": '[{"kind":"fill","role":"textbox",'
                                   '"name":"Email","value":"same-value"}]'}
        result = probe_profiles.match_transition_inputs(
            ParamsVault(), transition, {"ref_stem": "POST_items", "payload_json": "{}"},
        )
        self.assertEqual(result[0]["status"], "matched")
        self.assertEqual(result[0]["name"], "email")
        self.assertIn("name/label plus corroborating value", result[0]["reason"])

    def test_qty_matches_compound_entered_quantity_parameter(self):
        class ParamsVault:
            def get_params_for_api(self, _ref, role=None):
                return [{"name": "addtocart_31.EnteredQuantity"}]

        transition = {"path_json": '[{"kind":"fill","role":"textbox",'
                                   '"name":"Qty:","value":"1"}]'}
        result = probe_profiles.match_transition_inputs(
            ParamsVault(), transition,
            {"ref_stem": "POST_addproducttocart_details_31_1", "payload_json": "{}"},
        )
        self.assertEqual(result[0]["status"], "matched")
        self.assertEqual(result[0]["name"], "addtocart_31.EnteredQuantity")
        self.assertEqual(result[0]["confidence"], 0.86)

    def test_scope_drops_out_of_page_api_candidates(self):
        class Conn:
            def execute(self, *_args):
                return SimpleNamespace(fetchall=lambda: [])
        vault = SimpleNamespace(conn=Conn())
        transition = {"from_url": "https://shop.test/register", "path_json": "[]",
                      "triggered_apis_json": "[]"}
        with patch.object(testcase_gen, "_capture_api_submissions", return_value=[
            {"ref_stem": "POST_checkout_billing", "page_url": "https://shop.test/checkout",
             "submit_label": "Continue"},
            {"ref_stem": "POST_checkout_shipping", "page_url": "https://shop.test/checkout",
             "submit_label": "Continue"},
        ]):
            scoped, dropped = testcase_gen._transition_candidate_apis(
                vault, "s", transition,
                {"POST_checkout_billing": {}, "POST_checkout_shipping": {}},
            )
        self.assertEqual(scoped, {})
        self.assertEqual(dropped, 2)

    def test_capture_pairing_keeps_fields_through_focus_taps(self):
        self.vault.start_session("capture-session-focus", "https://example.test/form")
        self.vault.record_step("capture-session-focus", {
            "step_type": "input", "label": "Cardholder name", "value": "[REDACTED]",
            "page_url": "https://example.test/form",
        })
        self.vault.record_step("capture-session-focus", {
            "step_type": "tap", "label": "Card number", "page_url": "https://example.test/form",
            "action_id": "focus-card-number",
        })
        self.vault.record_step("capture-session-focus", {
            "step_type": "tap", "label": "Continue", "page_url": "https://example.test/form",
            "action_id": "submit-form",
        })
        self.vault.conn.execute(
            "INSERT INTO session_api_events (session_id,seq,ref_stem,action_id) VALUES (?,?,?,?)",
            ("capture-session-focus", 1, "POST_save", "submit-form"),
        )
        self.vault.conn.commit()
        paired = testcase_gen._capture_step_transitions(self.vault, "capture-session-focus")
        self.assertEqual(len(paired), 1)
        import json
        self.assertEqual([a["name"] for a in json.loads(paired[0]["path_json"])],
                         ["Cardholder name"])

    def test_response_classifier_rejects_200_error_body_and_accepts_captured_match(self):
        error = probe_profiles.classify_probe_response(
            200, '{"success":true,"message":"Saved"}',
            200, '{"success":false,"message":"Invalid email"}',
        )
        message_error = probe_profiles.classify_probe_response(
            200, '{"success":true,"message":"Saved"}',
            200, '{"message":"Invalid email"}',
        )
        same = probe_profiles.classify_probe_response(
            200, '{"success":true,"message":"Saved"}',
            200, '{"success":true,"message":"Saved"}',
        )
        self.assertEqual(error["outcome"], "rejected")
        self.assertEqual(message_error["outcome"], "rejected")
        self.assertEqual(same["outcome"], "accepted")

    def test_response_classifier_ignores_failure_word_already_in_good_control(self):
        captured = '{"message":"error: callback registered"}'
        same_control = probe_profiles.classify_probe_response(200, captured, 200, captured)
        new_failure = probe_profiles.classify_probe_response(
            200, captured, 200, '{"message":"Invalid search term"}',
        )
        self.assertEqual(same_control["outcome"], "accepted")
        self.assertEqual(new_failure["outcome"], "rejected")
        self.assertIn("invalid", new_failure["failure_indicators"])

    def test_unparseable_capture_limit_body_is_inconclusive_without_positive_signal(self):
        baseline = '{"success":true,"message":"Saved"}'
        truncated = '{"partial":"' + ("x" * (probe_profiles.MAX_BODY_CHARS - 12))
        result = probe_profiles.classify_probe_response(200, baseline, 200, truncated)
        self.assertEqual(result["outcome"], "inconclusive (truncated)")
        self.assertIn("capture limit", result["rule"])

    def test_unparseable_truncated_body_can_still_have_positive_evidence(self):
        baseline = '{"success":true,"message":"Saved"}'
        success_prefix = '{"success":true,"partial":"'
        failure_prefix = '<div>Invalid search query</div>'
        partial_success = success_prefix + ("x" * (probe_profiles.MAX_BODY_CHARS - len(success_prefix)))
        partial_failure = failure_prefix + ("x" * (probe_profiles.MAX_BODY_CHARS - len(failure_prefix)))
        accepted = probe_profiles.classify_probe_response(200, baseline, 200, partial_success)
        rejected = probe_profiles.classify_probe_response(200, baseline, 200, partial_failure)
        self.assertEqual(accepted["outcome"], "accepted")
        self.assertEqual(rejected["outcome"], "rejected")
        self.assertTrue(accepted.get("rule"))
        self.assertTrue(rejected.get("rule"))

    def test_non_json_classifier_ignores_script_and_style_text(self):
        captured = ('<div>Shipping details</div><script>error: function () { '
                    "alert('Failed to retrieve states.') }</script>"
                    '<style>.error { color: red }</style>')
        same = probe_profiles.classify_probe_response(200, captured, 200, captured)
        changed = probe_profiles.classify_probe_response(
            200, captured, 200, '<div>Invalid postal code</div><script>error</script>',
        )
        self.assertEqual(same["outcome"], "accepted")
        self.assertEqual(changed["outcome"], "rejected")
        self.assertIn("invalid", changed["failure_indicators"])

    def test_baseline_and_control_probe_logs_keep_more_than_document_limit(self):
        long_body = "x" * (probe_profiles.MAX_BODY_CHARS + 100)
        self.vault.save_field_probe("s::r", "POST_items", "r:1", "", "baseline",
                                    {}, "accepted", 200, "", response_body=long_body)
        self.vault.save_field_probe("s::r", "UI_search", "ui:1", "Search", "control",
                                    "shirt", "accepted with results", None, "",
                                    response_body=long_body)
        self.vault.save_field_probe("s::r", "POST_items", "r:1", "name", "short",
                                    "x", "rejected", 400, "invalid", response_body=long_body)
        rows = self.vault.get_field_probe_log("s::r")
        self.assertEqual(len(rows[0]["response_body"]), len(long_body))
        self.assertEqual(len(rows[1]["response_body"]), len(long_body))
        self.assertEqual(len(rows[2]["response_body"]), probe_profiles.MAX_BODY_CHARS)

    def test_session_control_response_uses_extended_capture_limit(self):
        body = "x" * (probe_profiles.MAX_BODY_CHARS + 200)
        self.vault.start_session("large-control", "https://example.test/")
        self.vault.ingest({
            "method": "POST", "url": "https://example.test/save",
            "request_headers": {"content-type": "application/json"},
            "response_headers": {"content-type": "application/json"},
            "payload": "{}", "response_body": body,
            "entry_point": "test", "ts": 1.0,
        }, session_id="large-control")
        stored = self.vault.conn.execute(
            "SELECT response_body FROM session_api_events WHERE session_id=?", ("large-control",),
        ).fetchone()[0]
        api_summary = self.vault.conn.execute(
            "SELECT response_json FROM apis WHERE method='POST' AND url_path='/save'",
        ).fetchone()[0]
        self.assertEqual(len(stored), len(body))
        self.assertLessEqual(len(stored), MAX_BASELINE_RESPONSE_CHARS)
        self.assertEqual(len(api_summary), probe_profiles.MAX_BODY_CHARS)

    def test_capture_ui_field_name_never_uses_entered_value(self):
        self.vault.start_session("capture-safe-name", "https://example.test/search")
        self.vault.record_step("capture-safe-name", {
            "step_type": "input", "label": "shirt", "value": "shirt",
            "selector": "input#q", "page_url": "https://example.test/search",
        })
        import json
        transition = testcase_gen._capture_input_transitions(self.vault, "capture-safe-name")[0]
        action = json.loads(transition["path_json"])[0]
        self.assertEqual(action["name"], "input#q")

    def test_capture_ui_grouping_merges_route_variants_by_form_and_selector(self):
        first = {"seq": "ui-1", "submit_label": "Add to cart"}
        second = {"seq": "ui-2", "submit_label": "Add to cart"}
        a = {"selector": "input#qty", "kind": "fill"}
        b = {"selector": "input#qty", "kind": "fill"}
        self.assertEqual(
            testcase_gen._field_group_scope(first, a, "https://shop.test/item/one", "Qty"),
            testcase_gen._field_group_scope(second, b, "https://shop.test/item/two", "Qty"),
        )
        self.assertEqual(
            testcase_gen._field_group_scope({"seq": 3}, {"kind": "fill"},
                                            "https://shop.test/item/one", "Qty"),
            testcase_gen._field_group_scope({"seq": 85}, {"kind": "fill"},
                                            "https://shop.test/", "Qty:"),
        )

    def test_cross_api_name_collision_is_reported_before_probing(self):
        class ParamsVault:
            def get_params_for_api(self, ref, role=None):
                return [{"name": "email", "examples": ["same-value"]}]

        transition = {"path_json": '[{"kind":"fill","role":"textbox",'
                                   '"name":"Email","value":"same-value"}]'}
        collisions = testcase_gen._cross_api_mapping_collisions(
            ParamsVault(), transition,
            {"POST_register": {"ref_stem": "POST_register", "payload_json": "{}"},
             "POST_login": {"ref_stem": "POST_login", "payload_json": "{}"}},
        )
        self.assertEqual(collisions, {0: [("POST_login", "email"), ("POST_register", "email")]})

    def test_debug_field_identity_deduplicates_repeated_named_fields(self):
        first = {"kind": "fill", "role": "textbox", "name": "Email", "value": "first"}
        later = {"kind": "fill", "role": "textbox", "name": "Email", "value": "later"}
        self.assertEqual(testcase_gen._debug_field_identity(first),
                         testcase_gen._debug_field_identity(later))

    def test_profile_groups_same_outcome_and_keeps_log_references(self):
        for value in ("", "x"):
            self.vault.save_field_probe("s::r", "POST_items", "r:1", "name", "empty",
                                        value, "rejected", 400, "invalid value")
        logs = self.vault.get_field_probe_log("s::r")
        profiles = probe_profiles.build_field_profiles(logs)
        profile = profiles[("POST_items", "name")]
        self.assertEqual(profile["observations"], 2)
        self.assertEqual(len(profile["outcomes"]), 1)
        self.assertEqual(profile["outcomes"][0]["confidence"], 1.0)
        self.assertEqual(len(profile["outcomes"][0]["raw_log_ids"]), 2)

    def test_profile_retains_low_confidence_page_form_match(self):
        profiles = probe_profiles.build_field_profiles([{
            "id": 19, "api_ref": "POST_addcart", "field_name": "Quantity",
            "outcome": "rejected", "response_status": 200,
            "match_confidence": 0.65, "source_of_truth": "recorded UI hint",
        }])
        self.assertEqual(profiles[("POST_addcart", "Quantity")]["match_confidence"], 0.65)

    def test_dry_run_planner_lists_scoped_register_field_without_request(self):
        class Conn:
            def execute(self, *_args):
                return SimpleNamespace(fetchall=lambda: [])
        vault = SimpleNamespace(
            conn=Conn(), vault_path=Path(self.temp.name),
            get_field_probe_log=lambda *_args: [],
            get_params_for_api=lambda *_args, **_kwargs: [],
        )
        transition = {
            "seq": 7, "from_url": "https://example.test/register",
            "path_json": '[{"kind":"fill","role":"textbox","name":"Email:","value":"a@example.test"}]',
            "action_kind": "", "action_name": "", "action_value": "",
            "triggered_apis_json": "[]",
        }
        api = {"ref_stem": "POST_checkout", "captured_response_status": 200,
               "captured_response_body": '{"success":true}', "method": "POST"}
        output = __import__("io").StringIO()
        with patch.object(testcase_gen, "_load_persisted_live_run", return_value=("run", [transition])), \
             patch.object(testcase_gen, "_session_probe_apis", return_value={"POST_checkout": api}), \
             patch.object(testcase_gen, "_capture_step_transitions", return_value=[]), \
             patch.object(testcase_gen, "_capture_input_transitions", return_value=[]), \
             patch.object(probe_profiles, "_send") as send, \
             patch("sys.stdout", output):
            testcase_gen.dry_run_probe_plan(vault, "session", "run")
        send.assert_not_called()
        self.assertIn("no browser, no requests", output.getvalue())
        self.assertIn("rule=", output.getvalue())
        self.assertIn("Register form preview", output.getvalue())
        self.assertIn("unique_registration_value", output.getvalue())
        self.assertIn("out-of-page API candidate(s) dropped", output.getvalue())

    def test_off_topic_or_similar_chunk_is_rejected_by_field_gate(self):
        field = {"name": "giftcard_2.RecipientEmail", "label": "Recipient's Email"}
        self.assertFalse(probe_profiles.chunk_concerns_field(
            {"title": "Add to cart", "text": "The cart quantity is updated after submission."}, field))
        self.assertTrue(probe_profiles.chunk_concerns_field(
            {"title": "Gift card", "text": "Recipient's Email is required for delivery."}, field))

    def test_weak_graph_neighbor_is_rejected(self):
        index_dir = Path(self.temp.name) / ".chat_index"
        index_dir.mkdir()
        (index_dir / "index.sqlite").touch()
        class FakeDB:
            def __init__(self, path):
                self.conn = object()
            def close(self):
                pass
        class FakeEngine:
            def __init__(self, db):
                pass
            def search(self, query, limit=8):
                return [{"id": 1, "source_path": "APIs/items.md", "ord": 0, "score": 0}]
            def expand_context(self, seeds, **kwargs):
                return [{"id": 2, "source_path": "BRDs/entries/email.md", "title": "Email",
                         "text": "Recipient email must use a valid format.", "origin": "expansion",
                         "edge_kind": "shared_entity", "edge_weight": .3, "final_score": .1}]
        field = {"name": "email", "label": "Recipient Email"}
        with patch("capture_app.chat.db.ChatDB", FakeDB), \
             patch("capture_app.chat.retriever.SearchEngine", FakeEngine):
            result = probe_profiles.find_rule_chunks(self.temp.name, {"ref_stem": "POST_items"}, field)
        self.assertEqual(result, [])

    def test_uncited_extraction_candidate_is_discarded(self):
        class FakeCache:
            def chat_json(self, *args, **kwargs):
                return {"POST_items::name": [{"field": "name", "constraint_kind": "max_length",
                    "value": 12, "source_node_ids": ["unprovided.md"]}]}
        result = probe_profiles.extract_written_rules([{
            "api_ref": "POST_items", "field": {"name": "name"},
            "chunks": [{"node_ids": ["BRDs/entries/a.md"], "text": "Name maximum length 12."}],
        }], FakeCache())
        self.assertEqual(result, {})

    def test_rule_prompt_redacts_sensitive_values(self):
        class FakeCache:
            prompt = ""
            def chat_json(self, system, user, **kwargs):
                self.prompt = system + user
                return {}
        cache = FakeCache()
        probe_profiles.extract_written_rules([{
            "api_ref": "POST_items", "field": {"name": "name"},
            "chunks": [{"node_ids": ["BRDs/entries/a.md"],
                        "text": "A captured contact was person@example.com; name length max 12."}],
        }], cache)
        self.assertNotIn("person@example.com", cache.prompt)

    def test_resume_never_exceeds_ten_probes_per_field(self):
        generated = [(f"kind{i}", i) for i in range(12)]
        self.api["captured_response_body"] = "ok"
        with patch.object(probe_profiles, "_candidate_probes", return_value=generated), \
             patch.object(probe_profiles, "_send", side_effect=[(200, "ok"), (400, "no")] * 20), \
             patch.object(probe_profiles.time, "sleep"):
            probe_profiles.probe_api(
                self.vault, "s::r", self.api, "r:1", [{"name": "name", "baseline": "known"}],
                confirm_state_change=True,
            )
            probe_profiles.probe_api(
                self.vault, "s::r", self.api, "r:1", [{"name": "name", "baseline": "known"}],
                confirm_state_change=True,
            )
        field_rows = [row for row in self.vault.get_field_probe_log("s::r")
                      if row["field_name"] == "name"]
        self.assertEqual(len(field_rows), 10)

    def test_field_stops_after_three_agreeing_probe_outcomes(self):
        with patch.object(probe_profiles, "_candidate_probes",
                          return_value=[("a", 1), ("b", 2), ("c", 3), ("d", 4)]), \
             patch.object(probe_profiles, "_send", side_effect=[(200, "baseline"), (400, "rejected"),
                                                                  (400, "rejected"), (400, "rejected")]) as send, \
             patch.object(probe_profiles.time, "sleep"):
            probe_profiles.probe_api(
                self.vault, "s::r", self.api, "r:1", [{"name": "name", "baseline": "known"}],
                confirm_state_change=True,
            )
        self.assertEqual(send.call_count, 4)  # baseline plus three agreeing probes

    def test_query_parameters_are_applied_to_get_url(self):
        api = {"method": "GET", "full_url": "https://example.test/search?old=1",
               "headers_json": "{}", "payload_json": "{}"}
        class FakeResponse:
            status_code = 200
            text = "ok"
        class FakeClient:
            last_url = ""
            def __init__(self, **kwargs):
                pass
            def __enter__(self):
                return self
            def __exit__(self, *args):
                return False
            def request(self, method, url, headers=None, **kwargs):
                FakeClient.last_url = url
                return FakeResponse()
        with patch.object(probe_profiles.httpx, "Client", FakeClient):
            probe_profiles._send(api, {"term": "cat"})
        self.assertIn("term=cat", FakeClient.last_url)
        self.assertIn("old=1", FakeClient.last_url)

    def test_get_variant_stays_in_query_even_with_captured_content_type(self):
        api = {"method": "GET", "full_url": "https://example.test/search?term=old",
               "headers_json": '{"Content-Type":"application/json"}', "payload_json": "{}"}
        class FakeResponse:
            status_code = 200
            text = "ok"
        class FakeClient:
            kwargs = None
            url = ""
            def __init__(self, **kwargs): pass
            def __enter__(self): return self
            def __exit__(self, *args): return False
            def request(self, method, url, headers=None, **kwargs):
                FakeClient.url, FakeClient.kwargs = url, kwargs
                return FakeResponse()
        with patch.object(probe_profiles.httpx, "Client", FakeClient):
            probe_profiles._send(api, {"term": "new value"})
        from urllib.parse import parse_qs, urlsplit
        self.assertEqual(parse_qs(urlsplit(FakeClient.url).query), {"term": ["new value"]})
        self.assertEqual(FakeClient.kwargs, {})

    def test_capture_inputs_pair_to_next_api_firing_tap_until_navigation(self):
        import json
        self.vault.start_session("capture-session", "https://example.test/form")
        self.vault.record_step("capture-session", {
            "step_type": "input", "label": "Email", "value": "a@example.test",
            "page_url": "https://example.test/form", "action_id": "input-1",
        })
        self.vault.record_step("capture-session", {
            "step_type": "input", "label": "Password", "value": "secret",
            "input_type": "password", "page_url": "https://example.test/form", "action_id": "input-2",
        })
        self.vault.record_step("capture-session", {
            "step_type": "tap", "label": "Sign in", "page_url": "https://example.test/form",
            "action_id": "submit-1",
        })
        self.vault.conn.execute(
            "INSERT INTO session_api_events (session_id,seq,ref_stem,action_id) VALUES (?,?,?,?)",
            ("capture-session", 1, "POST_login", "submit-1"),
        )
        self.vault.record_step("capture-session", {
            "step_type": "navigate", "page_url": "https://example.test/next", "action_id": "nav-1",
        })
        self.vault.record_step("capture-session", {
            "step_type": "input", "label": "Search", "value": "shoe",
            "page_url": "https://example.test/next", "action_id": "input-3",
        })
        self.vault.record_step("capture-session", {
            "step_type": "tap", "label": "Search", "page_url": "https://example.test/next",
            "action_id": "submit-2",
        })
        self.vault.conn.execute(
            "INSERT INTO session_api_events (session_id,seq,ref_stem,action_id) VALUES (?,?,?,?)",
            ("capture-session", 2, "GET_search", "submit-2"),
        )
        self.vault.conn.commit()
        paired = testcase_gen._capture_step_transitions(self.vault, "capture-session")
        self.assertEqual(len(paired), 2)
        self.assertEqual([a["name"] for a in json.loads(paired[0]["path_json"])], ["Email", "Password"])
        self.assertEqual(json.loads(paired[0]["triggered_apis_json"]), [{"ref_stem": "POST_login"}])
        self.assertEqual([a["name"] for a in json.loads(paired[1]["path_json"])], ["Search"])

    def test_enrich_get_variant_replaces_matching_query_value(self):
        url, method, _headers, json_body, form_body = enrich.apply_variant_to_request(
            {"method": "GET", "full_url": "https://example.test/search?term=old&page=1",
             "headers_json": "{}"},
            {"term": "new value"},
        )
        from urllib.parse import parse_qs, urlsplit
        self.assertEqual(method, "GET")
        self.assertEqual(parse_qs(urlsplit(url).query), {"term": ["new value"], "page": ["1"]})
        self.assertIsNone(json_body)
        self.assertIsNone(form_body)

    def test_sensitive_probe_values_and_response_fields_are_redacted_on_save(self):
        self.vault.save_field_probe(
            "s::r", "POST_items", "r:1", "password", "wrong_type", "secret-value",
            "rejected", 400, "bad password secret-value",
            response_body='{"password":"secret-value","message":"invalid"}',
        )
        row = self.vault.get_field_probe_log("s::r")[0]
        self.assertEqual(row["probe_value"], "[TEST VALUE REQUIRED]")
        self.assertEqual(row["response_message"], "")
        self.assertEqual(row["response_body"], "")

    def test_negative_case_uses_observed_rejection_and_changes_one_field(self):
        api = {"id": 1, "ref_stem": "POST_items", "method": "POST",
               "full_url": "https://example.test/items", "headers_json": "{}",
               "payload_json": '{"name":"known","quantity":5}'}
        rejection_id = self.vault.save_field_probe(
            "s::r", "POST_items", "r:1", "name", "wrong_type", "invalid-name",
            "rejected", 422, "Name is invalid", response_body='{"message":"Name is invalid"}',
        )
        params = [
            {"name": "name", "examples": ["known"], "resource_key": "example:items"},
            {"name": "quantity", "examples": [5], "resource_key": "example:items"},
        ]
        fake_vault = SimpleNamespace(
            get_field_probe_log=self.vault.get_field_probe_log,
            list_apis_with_params=lambda: [api],
            get_params_for_api=lambda ref, role=None: params,
        )
        row = {
            "Test Case Name": "Create item", "Sub-module": "Create item",
            "Data": "name: known; quantity: 5", "Steps": "", "Expected Result": "Created",
            "_transition_key": "r:1",
            "_transition_actions": [
                {"kind": "fill", "role": "textbox", "name": "name", "value": "known"},
                {"kind": "fill", "role": "textbox", "name": "quantity", "value": "5"},
            ],
            "_transition_apis": [{"method": "POST", "url": "/items", "ref_stem": "POST_items"}],
        }
        generated = testcase_gen._profile_case_rows(fake_vault, [row], "s::r")
        negative = next(item for item in generated if item["Category"] == "negative")
        self.assertEqual(negative["Data"], "name: invalid-name")
        self.assertIn("invalid-name", negative["Steps"])
        self.assertIn("Enter 5 in textbox quantity", negative["Steps"])
        self.assertIn(f"#{rejection_id}", negative["Technical Evidence"])
        self.assertNotIn("HTTP", negative["Expected Result"])

    def test_unknown_field_mapping_does_not_invent_a_case(self):
        api = {"id": 1, "ref_stem": "POST_items", "method": "POST",
               "full_url": "https://example.test/items", "headers_json": "{}",
               "payload_json": '{"name":"known"}'}
        self.vault.save_field_probe("s::r", "POST_items", "r:1", "name", "wrong_type",
                                    "bad", "rejected", 400, "invalid")
        fake_vault = SimpleNamespace(
            get_field_probe_log=self.vault.get_field_probe_log,
            list_apis_with_params=lambda: [api],
            get_params_for_api=lambda ref, role=None: [
                {"name": "name", "examples": ["known"], "resource_key": "example:items"}],
        )
        row = {"Test Case Name": "Create item", "_transition_key": "r:1",
               "_transition_actions": [{"kind": "fill", "name": "Account", "value": "other"}],
               "_transition_apis": [{"method": "POST", "url": "/items", "ref_stem": "POST_items"}]}
        self.assertEqual(testcase_gen._profile_case_rows(fake_vault, [row], "s::r"), [])

    def test_wording_failure_keeps_deterministic_fallback(self):
        rows = [{"Test Case Name": "Create item - negative name",
                 "Expected Result": "The system rejects the entered name value.",
                 "Technical Evidence": "Observed status 422; raw probe log #1"}]
        fake_vault = SimpleNamespace(vault_path=Path(self.temp.name))
        with patch.object(testcase_gen._ResponseCache, "chat_json",
                          side_effect=testcase_gen.OpenRouterError("offline")):
            testcase_gen._word_profile_expectations(rows, fake_vault)
        self.assertEqual(rows[0]["Expected Result"], "The system rejects the entered name value.")

    def test_enhanced_workbook_adds_evidence_after_tester_columns(self):
        from openpyxl import load_workbook
        row = {"Module": "Items", "Sub-module": "Invalid name", "Test Case Name": "Reject invalid name",
               "Precondition": "None", "Data": "name: bad", "Steps": "1. Enter bad in Name",
               "Expected Result": "The system rejects the entered name value.",
               "Category": "negative", "Technical Evidence": "Observed status 422; raw probe log #1"}
        output = testcase_gen.write_excel([row], Path(self.temp.name) / "enhanced.xlsx", enhanced=True)
        sheet = load_workbook(output, data_only=True)["Test Cases"]
        headers = [sheet.cell(5, column).value for column in range(1, 13)]
        self.assertEqual(headers[8:11], ["Actual Result", "Status", "Evidence/Note"])
        self.assertEqual(headers[11], "Technical Evidence")
        self.assertEqual(sheet.cell(6, 10).value, "Untested")
        self.assertEqual(sheet.cell(6, 12).value, row["Technical Evidence"])

    def test_rule_probe_disagreement_is_written_to_review_output(self):
        rule_id = self.vault.save_field_written_rule(
            "s::r", "POST_items", "name", "maximum", 10, "", ["BRDs/entries/names.md"],
            "", strength="written rule candidate")
        self.vault.save_field_probe("s::r", "POST_items", "r:1", "name", "rule_limit", 10,
                                    "rejected", 422, "too long")
        probe_id = self.vault.save_field_probe("s::r", "POST_items", "r:1", "name", "rule_violation", 11,
                                              "accepted", 200, "ok")
        reviews = testcase_gen._rule_review_rows(self.vault, "s::r")
        self.assertEqual(len(reviews), 1)
        self.assertEqual(reviews[0]["state"], "contradicted")
        self.assertIn("BRDs/entries/names.md", reviews[0]["source_node_ids"])
        self.assertIn(f"#{probe_id}", reviews[0]["observed_evidence"])
        output = testcase_gen.write_excel([], Path(self.temp.name) / "review.xlsx", enhanced=True,
                                          graph_rules=True, rule_reviews=reviews)
        from openpyxl import load_workbook
        workbook = load_workbook(output, data_only=True)
        self.assertEqual(workbook["Rule Review"].cell(2, 7).value, "contradicted")
        self.assertEqual(self.vault.get_field_written_rules("s::r")[0]["id"], rule_id)

    def test_only_a_confirmed_written_limit_creates_boundary_case(self):
        api = {"id": 1, "ref_stem": "POST_items", "method": "POST",
               "full_url": "https://example.test/items", "headers_json": "{}",
               "payload_json": '{"name":"known"}'}
        params = [{"name": "name", "examples": ["known"], "resource_key": "example:items"}]
        self.vault.save_field_written_rule("s::r", "POST_items", "name", "max_length", 4,
                                           "", ["BRDs/entries/names.md"], "")
        self.vault.save_field_probe("s::r", "POST_items", "r:1", "name", "rule_limit", "xxxx",
                                    "accepted", 200, "ok")
        self.vault.save_field_probe("s::r", "POST_items", "r:1", "name", "rule_violation", "xxxxx",
                                    "rejected", 422, "too long")
        fake_vault = SimpleNamespace(
            get_field_probe_log=self.vault.get_field_probe_log,
            get_field_written_rules=self.vault.get_field_written_rules,
            update_field_written_rule=self.vault.update_field_written_rule,
            list_apis_with_params=lambda: [api],
            get_params_for_api=lambda ref, role=None: params,
        )
        row = {"Test Case Name": "Create item", "_transition_key": "r:1",
               "_transition_actions": [{"kind": "fill", "role": "textbox", "name": "name", "value": "known"}],
               "_transition_apis": [{"method": "POST", "url": "/items", "ref_stem": "POST_items"}]}
        cases = testcase_gen._profile_case_rows(fake_vault, [row], "s::r", graph_rules=True)
        boundary = next(item for item in cases if item["Category"] == "boundary")
        self.assertEqual(boundary["Data"], "name: xxxx")
        self.assertEqual(boundary["Basis"], "Rule + Observed")
        self.assertIn("confirmed", boundary["Technical Evidence"])

    def test_graph_flag_off_keeps_probe_only_workbook_columns(self):
        from openpyxl import load_workbook
        row = {"Module": "Items", "Test Case Name": "Baseline", "Steps": "Submit",
               "Expected Result": "Saved", "Technical Evidence": "Observed"}
        path = testcase_gen.write_excel([row], Path(self.temp.name) / "probe-only.xlsx", enhanced=True,
                                        graph_rules=False)
        sheet = load_workbook(path, data_only=True)["Test Cases"]
        headers = [sheet.cell(5, column).value for column in range(1, 13)]
        self.assertEqual(headers[-1], "Technical Evidence")
        self.assertEqual(sheet.max_column, 12)
        self.assertEqual(load_workbook(path, data_only=True).sheetnames, ["Test Cases"])


if __name__ == "__main__":
    unittest.main()
