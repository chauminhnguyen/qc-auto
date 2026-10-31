import json
import sqlite3
import tempfile
import unittest
from contextlib import redirect_stdout
from io import StringIO
from pathlib import Path
from types import SimpleNamespace
from unittest.mock import patch

from capture_app.integrations.openrouter import OpenRouterError
from capture_app.workflows import testcase_gen


class IsolatedGenerationTestCase(unittest.TestCase):
    """Keep persistent model caches and resumable state out of unit tests."""

    def setUp(self):
        self.provider_patch = patch.object(testcase_gen, "LLM_PROVIDER", "openrouter")
        self.provider_patch.start()
        self.addCleanup(self.provider_patch.stop)
        temp_root = Path(tempfile.mkdtemp(prefix="testcase-gen-tests-"))
        self.addCleanup(lambda: __import__("shutil").rmtree(temp_root, ignore_errors=True))
        original_cache = testcase_gen._ResponseCache
        self.cache_patch = patch.object(
            testcase_gen, "_ResponseCache",
            side_effect=lambda path, enabled=True, max_calls=testcase_gen.MAX_GENERATION_CALLS:
                original_cache(path, enabled=False, max_calls=max_calls),
        )
        self.cache_patch.start()
        self.addCleanup(self.cache_patch.stop)

        original_load = testcase_gen._load_brd_progress
        original_save = testcase_gen._save_brd_progress
        self.load_patch = patch.object(
            testcase_gen, "_load_brd_progress",
            side_effect=lambda path, fingerprint: original_load(temp_root / Path(path).name, fingerprint),
        )
        self.save_patch = patch.object(
            testcase_gen, "_save_brd_progress",
            side_effect=lambda path, progress: original_save(temp_root / Path(path).name, progress),
        )
        self.load_patch.start()
        self.save_patch.start()
        self.addCleanup(self.save_patch.stop)
        self.addCleanup(self.load_patch.stop)


def endpoint(ref_stem: str) -> dict:
    return {
        "ref_stem": ref_stem,
        "endpoint": f"POST /{ref_stem}",
        "friendly_name": ref_stem,
        "source": "captured",
        "preconditions": [],
        "precondition_flow_files": {},
        "flow_file": None,
        "destructive": False,
        "mutating": True,
        "example_request": "field=value",
        "example_response": "ok",
        "request_fields_with_real_examples": {"field": ["value", "another"]},
        "real_status": 200,
    }


def group(ref_stem: str, name: str | None = None) -> dict:
    return {
        "ref_stem": ref_stem,
        "cases": [{
            "category": "invalid-input",
            "name": name or f"{ref_stem} invalid input",
            "precondition": "None",
            "data": "field=invalid",
            "data_plain": "Set the field to the invalid value shown.",
            "steps": ["Open the feature.", "Enter field=invalid.", "Submit the request."],
            "expected_result": "The request is rejected",
        }],
        "happy_path_steps": ["Open the feature.", "Use the captured request.", "Check the result."],
        "happy_path_data": "field=value",
        "happy_path_data_plain": "Field: value",
        "happy_path_precondition_plain": "None",
        "happy_path_expected_result_plain": "A confirmation message appears.",
    }


class GenerateBatchTests(IsolatedGenerationTestCase):
    def test_anthropic_cache_call_omits_unsupported_temperature(self):
        cache = testcase_gen._ResponseCache(Path("unused-cache.json"), enabled=False)
        with patch.object(testcase_gen, "LLM_PROVIDER", "anthropic"), \
             patch.object(testcase_gen, "messages_json", return_value={"ok": True}) as messages:
            result = cache.chat_json("system", "user", model=None, max_tokens=100)

        self.assertEqual(result, {"ok": True})
        self.assertIsNone(messages.call_args.kwargs["temperature"])

    @staticmethod
    def write_knowledge_note(root: Path, *, note_id="K_subscribe", target="POST_subscribenewsletter",
                             status="unverified"):
        note_dir = root / "Knowledge" / "Reviewed"
        note_dir.mkdir(parents=True, exist_ok=True)
        note_path = note_dir / f"{note_id}.md"
        metadata = {
            "type": "business-rule",
            "feature": "Subscribe Button",
            "title": "Subscribe button color requirement",
            "status": status,
            "confirmed_by": [],
            "applies_to": [f"[[APIs/demowebshop.tricentis.com/{target}]]"],
        }
        frontmatter = "\n".join(f"{key}: {json.dumps(value)}" for key, value in metadata.items())
        note_path.write_text(
            f"---\n{frontmatter}\n---\n\n# Subscribe button color requirement\n\n"
            "The subscribe button must be gray and black text\n\n"
            "## Source evidence\nthe subscribe button must be gray and black text\n\n"
            "## Open questions\n- Is this for the newsletter page?\n",
            encoding="utf-8",
        )
        return note_path

    def test_partial_batch_retries_only_missing_endpoint(self):
        with patch.object(testcase_gen, "chat_json", side_effect=[
            [group("first")],
            [group("second")],
        ]) as chat:
            rows = testcase_gen.generate_batch([endpoint("first"), endpoint("second")])

        self.assertEqual({row["Endpoint"] for row in rows}, {"POST /first", "POST /second"})
        self.assertEqual(chat.call_count, 2)
        self.assertIn('"ref_stem": "second"', chat.call_args.args[1])
        self.assertNotIn('"ref_stem": "first"', chat.call_args.args[1])
        self.assertIn('"edge_case_eligible": true', chat.call_args.args[1])
        self.assertNotIn("precondition_flow_files", chat.call_args.args[1])
        self.assertTrue(chat.call_args.args[0].endswith("omit explanations."))
        self.assertIn("field=invalid", rows[0]["Data"])
        self.assertIn("1. Open the feature.", rows[0]["Steps"])

    def test_truncated_array_keeps_complete_group_and_retries_missing(self):
        first_group = json.dumps(group("first"))
        truncated = f"[{first_group}, {{\"ref_stem\":\"second\",\"cases\":[{{\"category\":"
        with patch.object(testcase_gen, "chat_json", side_effect=[
            OpenRouterError("invalid json", raw_output=truncated),
            [group("second")],
        ]) as chat:
            rows = testcase_gen.generate_batch([endpoint("first"), endpoint("second")])

        self.assertEqual({row["Endpoint"] for row in rows}, {"POST /first", "POST /second"})
        self.assertEqual(chat.call_count, 2)
        self.assertIn('"ref_stem": "second"', chat.call_args.args[1])

    def test_missing_happy_path_steps_retry_without_duplicating_cases(self):
        incomplete = group("only")
        incomplete.pop("happy_path_steps")
        with patch.object(testcase_gen, "chat_json", side_effect=[
            [incomplete],
            [group("only")],
        ]) as chat:
            rows = testcase_gen.generate_batch([endpoint("only")])

        self.assertEqual(chat.call_count, 2)
        self.assertEqual(len(rows), 1)

    def test_single_endpoint_failure_gets_one_compact_retry(self):
        with patch.object(testcase_gen, "chat_json", side_effect=[
            OpenRouterError("invalid json", raw_output="[truncated"),
            [group("only")],
        ]) as chat:
            rows = testcase_gen.generate_batch([endpoint("only")], max_tokens=2400)

        self.assertEqual(len(rows), 1)
        self.assertEqual(chat.call_count, 2)
        self.assertEqual(chat.call_args.kwargs["max_tokens"], 2400)
        self.assertIn("at most 2 concise cases", chat.call_args.args[0])
        self.assertNotIn("another", chat.call_args.args[1])

    def test_duplicate_cases_are_not_written_twice(self):
        duplicate = group("only")
        with patch.object(testcase_gen, "chat_json", return_value=[duplicate, duplicate]):
            rows = testcase_gen.generate_batch([endpoint("only")])

        self.assertEqual(len(rows), 1)

    def test_same_name_with_distinct_data_is_preserved(self):
        first = group("only", name="Same scenario")
        second = group("only", name="Same scenario")
        second["cases"][0]["data"] = "field=other"
        with patch.object(testcase_gen, "chat_json", return_value=[first, second]):
            rows = testcase_gen.generate_batch([endpoint("only")])

        self.assertEqual(len(rows), 2)

    def test_compact_failure_is_bounded(self):
        with patch.object(testcase_gen, "chat_json", side_effect=OpenRouterError("invalid json")) as chat:
            rows = testcase_gen.generate_batch([endpoint("only")])

        self.assertEqual(rows, [])
        self.assertEqual(chat.call_count, 2)

    def test_row_metadata_is_stable_and_case_specific(self):
        ctx = endpoint("checkout/save")
        case = group("checkout/save")["cases"][0]
        first = testcase_gen._case_to_row(case, ctx)
        second = testcase_gen._case_to_row(case, ctx)
        changed = dict(case, data="field=other")
        third = testcase_gen._case_to_row(changed, ctx)

        self.assertEqual(first["ID"], second["ID"])
        self.assertNotEqual(first["ID"], third["ID"])
        self.assertEqual((first["Module"], first["Sub-module"]), ("Checkout", first["Test Case Name"]))
        self.assertEqual(first["Cycle"], "Regression, Integration")
        self.assertEqual(first["Actual Result"], "")
        self.assertEqual(first["Status/Note"], "")
        happy = testcase_gen._happy_path_row(ctx)
        self.assertEqual(happy["Data"], "")
        self.assertEqual(happy["Captured Request"], "field=value")
        self.assertIn("Blocked", happy["Review Status"])
        self.assertEqual(happy["Cycle"], "UAT, Regression, Integration")

    def test_model_output_without_concrete_data_is_blocked_not_executable(self):
        case = group("only")["cases"][0]
        case["data"] = ""
        row = testcase_gen._case_to_row(case, endpoint("only"))
        self.assertEqual(row["Data"], "")
        self.assertEqual(row["Review Status"], "Blocked – needs clarification")
        self.assertIn("concrete test data", row["Review Note"])

    def test_all_contexts_get_steps_but_discoveries_get_no_edge_cases(self):
        captured = endpoint("eligible")
        noneligible = endpoint("read-only")
        noneligible.update({
            "endpoint": "GET /read-only",
            "mutating": False,
            "example_request": "",
            "request_fields_with_real_examples": {},
        })
        discovery = {
            "ref_stem": "discovered",
            "endpoint": "POST /discovered",
            "friendly_name": "Discovered Action",
            "source": "discovery",
            "triggering_action": "click 'Continue'",
            "real_response_status": 200,
            "precondition_flow_files": {},
        }
        with patch.object(testcase_gen, "generate_supplemental_expected_results", side_effect=lambda rows, *_a, **_k: rows) as supplemental, \
             patch.object(testcase_gen, "build_contexts", return_value=[captured, noneligible, discovery]), \
             patch.object(testcase_gen, "chat_json", return_value=[
                 group("eligible"),
                 group("read-only"),
                 group("discovered"),
             ]) as chat:
            rows = testcase_gen.generate_all(None, brd_db_path=None)

        self.assertEqual(len(rows), 4)
        self.assertTrue(all("Suggested steps (verify before use)" not in row["Steps"] for row in rows))
        self.assertTrue(all(row["Review Status"] for row in rows))
        self.assertEqual(sum(row["Category"] != "happy-path" for row in rows), 1)
        captured_data = [row["Captured Request"] for row in rows if row["Category"] == "happy-path"]
        self.assertIn("field=value", captured_data)
        supplemental.assert_called_once()
        sent_contexts = json.loads(chat.call_args.args[1])
        self.assertEqual([ctx["edge_case_eligible"] for ctx in sent_contexts], [True, False, False])

    def test_flow_happy_path_uses_synthetic_readable_data_and_explicit_ui_outcome(self):
        ctx = endpoint("subscribenewsletter")
        ctx.update({
            "example_request": "",
            "example_request_needs_synthetic": True,
            "preconditions": ["POST /checkout/savebilling"],
            "_generated_happy_path_data": "email=new.user@example.test",
            "_generated_happy_path_data_plain": "Email address: new.user@example.test",
            "_generated_happy_path_precondition_plain": "The billing address is saved.",
            "_generated_happy_path_expected_result_plain": "A subscription confirmation appears.",
            "_generated_happy_steps": [
                "Open the newsletter signup form.",
                "Enter new.user@example.test and select Subscribe.",
            ],
        })
        row = testcase_gen._happy_path_row(ctx)

        self.assertEqual(row["Data"], "email=new.user@example.test")
        self.assertEqual(row["Data (Plain Language)"], "Email address: new.user@example.test")
        self.assertEqual(row["Captured Request"], "")
        self.assertEqual(row["Precondition (Plain Language)"], "The billing address is saved.")
        self.assertEqual(row["Expected Result"], "A subscription confirmation appears.")
        self.assertTrue(row["Steps"].startswith("1. "))
        self.assertNotIn("Suggested steps", row["Steps"])
        self.assertTrue(row["Review Status"].startswith("Generated"))

    def test_excel_is_one_sheet_with_status_summary_and_human_entry_fields(self):
        row = testcase_gen._case_to_row(
            group("cart/add")["cases"][0],
            endpoint("cart/add"),
        )
        row["_Ref Stem"] = "POST_cart/add"
        row["_Precondition Refs"] = ["POST_checkout/save"]
        row["Precondition"] = "POST /checkout/save has succeeded."
        prerequisite = {
            "Module": "Checkout",
            "Sub-module": "Save billing address",
            "Test Case Name": "Save billing address",
            "Precondition": "None",
            "Data": "Address: 1 Example Way",
            "Steps": "1. Open checkout.\n2. Save the billing address.",
            "Expected Result": "The saved billing address is shown in the checkout summary.",
            "Category": "happy-path",
            "Review Status": "Generated; verify before execution",
            "_Ref Stem": "POST_checkout/save",
            "_Precondition Refs": [],
        }
        smoke = {**prerequisite, "Module": "Catalog", "_Ref Stem": "GET_catalog", "_Precondition Refs": []}
        regression = {
            **row, "Module": "Catalog", "Test Case Name": "Regression case",
            "_Ref Stem": "GET_categories", "_Precondition Refs": [],
        }
        blocked = {**prerequisite, "Test Case Name": "Blocked case", "Review Status": "Blocked – needs clarification"}
        nfr = {"Output Sheet": "NFR", "Test Case Name": "Availability requirement"}
        with tempfile.TemporaryDirectory() as temp_dir:
            path = testcase_gen.write_excel(
                [row, prerequisite, smoke, regression, blocked, nfr], Path(temp_dir) / "testcases.xlsx"
            )
            from openpyxl import load_workbook

            workbook = load_workbook(path, data_only=False)
            self.assertEqual(workbook.sheetnames, ["Test Cases"])
            sheet = workbook["Test Cases"]
            self.assertEqual(
                [sheet.cell(2, column).value for column in range(1, 6)],
                ["Passed", "Failed", "Untested", "Pending", "Unknown"],
            )
            for column in range(1, 6):
                self.assertIn("COUNTIF", sheet.cell(3, column).value)
            headers = [sheet.cell(5, column).value for column in range(1, 12)]
            self.assertEqual(headers, [
                "Module", "Sub-module (Test Summary)", "Test Case Name", "Precondition", "Data",
                "Steps", "Expected Result", "Testing Type", "Actual Result", "Status", "Evidence/Note",
            ])
            values_by_name = {
                sheet.cell(row_num, 3).value: {
                    header: sheet.cell(row_num, column).value
                    for column, header in enumerate(headers, start=1)
                }
                for row_num in range(6, sheet.max_row + 1)
            }
            self.assertEqual(set(values_by_name), {
                row["Test Case Name"], "Save billing address", "Regression case",
            })
            self.assertEqual(values_by_name[row["Test Case Name"]]["Testing Type"], "Integration")
            self.assertEqual(values_by_name["Save billing address"]["Testing Type"], "Smoke")
            self.assertEqual(values_by_name["Regression case"]["Testing Type"], "Regression")
            for values in values_by_name.values():
                self.assertIn(values["Actual Result"], ("", None))
                self.assertIn(values["Evidence/Note"], ("", None))
                self.assertEqual(values["Status"], "Untested")
            self.assertIn("J6:J1048576", [str(item.sqref) for item in sheet.data_validations.dataValidation])
            self.assertEqual(len(sheet.tables), 1)
            self.assertTrue(sheet.tables["TestCases"].autoFilter.ref.startswith("A5:L"))
            self.assertEqual(sheet.freeze_panes, "A6")
            workbook.close()


    def test_batch_maps_out_of_order_results_and_skips_explicit_brd_oracle(self):
        rows = [
            {"ID": "first", "Test Case Name": "Explicit", "Steps": "1. Save",
             "BRD Evidence": "The application displays 'Saved' after saving."},
            {"ID": "second", "Test Case Name": "Ad hoc", "Steps": "1. Save",
             "BRD Evidence": ""},
        ]
        response = {"results": [
            {"line_id": "second", "has_explicit_brd_oracle": False,
             "generated_expected_result": "The saved item appears in the list."},
            {"line_id": "first", "has_explicit_brd_oracle": True,
             "generated_expected_result": ""},
        ]}
        cache = testcase_gen._ResponseCache(Path("unused-cache.json"), enabled=False)
        with patch.object(testcase_gen, "chat_json", return_value=response) as chat:
            testcase_gen.generate_supplemental_expected_results(rows, cache)

        self.assertEqual(rows[0].get("Generated Expected Result"), None)
        self.assertEqual(rows[1]["Generated Expected Result"], "The saved item appears in the list.")
        submitted = json.loads(chat.call_args.args[1])["rows"]
        self.assertEqual([item["line_id"] for item in submitted], ["first", "second"])

    def test_batch_rejects_unknown_duplicate_and_missing_line_ids(self):
        item = {"line_id": "row-6"}
        for results, message in (
            ([{"line_id": "other", "has_explicit_brd_oracle": False,
               "generated_expected_result": "A result."}], "Unknown"),
            ([{"line_id": "row-6", "has_explicit_brd_oracle": False,
               "generated_expected_result": "A result."}] * 2, "Duplicate"),
            ([], "omitted"),
        ):
            with self.subTest(message=message), self.assertRaisesRegex(ValueError, message):
                testcase_gen._validate_generated_expected_result_batch({"results": results}, [item])

    def test_helper_updates_only_empty_result_cells_and_preserves_original_workbook(self):
        from openpyxl import load_workbook

        with tempfile.TemporaryDirectory() as temp_dir:
            root = Path(temp_dir)
            source = testcase_gen.write_excel([
                {"Module": "Orders", "Test Case Name": "Needs result", "Steps": "1. Save",
                 "Expected Result": "Primary text stays."},
                {"Module": "Orders", "Test Case Name": "Already filled", "Steps": "1. Open",
                 "Expected Result": "Original expected.", "Generated Expected Result": "Existing supplemental."},
            ], root / "cases.xlsx")
            output = root / "completed.xlsx"
            response = {"results": [{
                "line_id": "Test Cases!6", "has_explicit_brd_oracle": False,
                "generated_expected_result": "The order appears in the list.",
            }]}
            with patch.object(testcase_gen, "chat_json", return_value=response) as chat:
                testcase_gen.fill_generated_expected_results_workbook(
                    source, output_path=output,
                    cache=testcase_gen._ResponseCache(root / "cache.json", enabled=False),
                )

            self.assertFalse(output == source)
            self.assertFalse(load_workbook(source, data_only=True)["Test Cases"]["L6"].value)
            book = load_workbook(output, data_only=True)
            sheet = book["Test Cases"]
            self.assertEqual(sheet["G6"].value, "Primary text stays.")
            self.assertEqual(sheet["L6"].value, "The order appears in the list.")
            self.assertEqual(sheet["L7"].value, "Existing supplemental.")
            self.assertIn("Test Cases!6", chat.call_args.args[1])
            book.close()

    def test_step_formatter_preserves_arbitrarily_many_long_steps(self):
        steps = [f"Open screen {index} and verify the displayed value: {index}." + " Detail" * 60
                 for index in range(9)]
        formatted = testcase_gen._format_steps(steps)

        self.assertEqual(len(formatted.splitlines()), 9)
        self.assertTrue(formatted.splitlines()[-1].startswith("9. "))
        self.assertGreater(len(formatted.splitlines()[0]), 300)

    def test_case_row_keeps_more_than_six_steps_and_detailed_expected_result(self):
        case = group("checkout/save")["cases"][0]
        case["steps"] = [f"Perform action {index}." for index in range(1, 10)]
        case["expected_result"] = "The expected screen is shown in sequence. " * 40
        case["expected_result_plain"] = case["expected_result"]
        row = testcase_gen._case_to_row(case, endpoint("checkout/save"))

        self.assertIsNotNone(row)
        self.assertEqual(len(row["Steps"].splitlines()), 9)
        self.assertEqual(row["Expected Result"], case["expected_result"].strip())

    def test_knowledge_note_is_read_and_matched_by_linked_api(self):
        with tempfile.TemporaryDirectory() as temp_dir:
            root = Path(temp_dir)
            path = self.write_knowledge_note(root)
            notes = testcase_gen._load_knowledge_notes(root)
        ctx = endpoint("subscribenewsletter")
        ctx["ref_stem"] = "POST_subscribenewsletter"

        self.assertEqual(len(notes), 1)
        self.assertEqual(notes[0]["knowledge_id"], "K_subscribe")
        self.assertEqual(notes[0]["status"], "unverified")
        self.assertIn("must be gray and black text", notes[0]["claim"])
        self.assertEqual(notes[0]["open_questions"], ["Is this for the newsletter page?"])
        self.assertEqual(notes[0]["source_path"], path.relative_to(root).as_posix())
        self.assertEqual(
            testcase_gen._knowledge_notes_for_context(ctx, notes),
            notes,
        )
        self.assertEqual(testcase_gen._knowledge_notes_for_context(endpoint("other"), notes), [])

    def test_open_question_knowledge_is_blocked_in_review_queue(self):
        ctx = endpoint("subscribenewsletter")
        ctx["ref_stem"] = "POST_subscribenewsletter"
        ctx["friendly_name"] = "Subscribe Newsletter"
        with tempfile.TemporaryDirectory() as temp_dir:
            root = Path(temp_dir)
            self.write_knowledge_note(root)
            vault = SimpleNamespace(vault_path=root)
            with patch.object(testcase_gen, "generate_supplemental_expected_results", side_effect=lambda rows, *_a, **_k: rows), \
                 patch.object(testcase_gen, "build_contexts", return_value=[ctx]), \
                 patch.object(testcase_gen, "chat_json", side_effect=[
                         [group("POST_subscribenewsletter")],
                     ]) as chat:
                    rows = testcase_gen.generate_all(vault, brd_db_path=None)

        self.assertEqual(chat.call_count, 1)
        self.assertEqual(len(rows), 3)
        review_row = next(row for row in rows if row.get("Output Sheet") == "Review Queue")
        self.assertEqual(review_row["Source Reference"], "K_subscribe")
        self.assertEqual(review_row["Review Status"], "Blocked – needs clarification")
        self.assertIn("Is this for the newsletter page?", review_row["Review Note"])
        with tempfile.TemporaryDirectory() as temp_dir:
            workbook_path = testcase_gen.write_excel(rows, Path(temp_dir) / "testcases.xlsx")
            from openpyxl import load_workbook

            workbook = load_workbook(workbook_path, data_only=True)
            self.assertEqual(workbook.sheetnames, ["Test Cases"])
            sheet = workbook["Test Cases"]
            self.assertEqual(sheet.cell(5, 1).value, "Module")
            self.assertEqual(sheet.max_row, 5)
            workbook.close()


class BrdCaseGenerationTests(IsolatedGenerationTestCase):
    @staticmethod
    def requirement(req_id="BR-01.3", statement="Each account email address is unique.",
                    source_file="BRD-01_Registration.md", entry_type="requirement",
                    priority="Must", status="approved"):
        return {
            "entry_id": f"entry-{source_file}-{req_id}",
            "requirement_id": req_id,
            "statement": statement,
            "priority": priority,
            "entry_type": entry_type,
            "source_path": f"D:\\BRDs\\{source_file}",
            "source_file": source_file,
            "source_display": source_file,
            "section": "Registration",
            "page": 2,
            "source_version": "1.0",
            "status": status,
            "open_questions": [],
        }

    @staticmethod
    def generated_item(trace_key="BRD-01_Registration.md/BR-01.3"):
        return {
            "traceability_key": trace_key,
            "classification": "testable",
            "module": "Registration",
            "sub_module": "Email uniqueness",
            "review_status": "Ready for review",
            "review_note": "",
            "endpoint_ref": None,
            "cases": [
                {
                    "category": "positive",
                    "name": "Register with an unused email address",
                    "precondition": "The registration page is open.",
                    "precondition_plain": "The registration page is open.",
                    "data": "Email: new.user@example.test",
                    "data_plain": "Use an email address not already registered.",
                    "steps": ["Enter the new email address.", "Complete required fields and submit."],
                    "expected_result": "The new account is created and the account page opens.",
                    "expected_result_plain": "The new account is created and the account page opens.",
                },
                {
                    "category": "negative",
                    "name": "Register with an email already in use",
                    "precondition": "An account already exists for existing.user@example.test.",
                    "precondition_plain": "An account already exists for the email address.",
                    "data": "Email: existing.user@example.test",
                    "data_plain": "Use the email address already registered.",
                    "steps": ["Enter the existing email address.", "Complete other required fields and submit."],
                    "expected_result": "An inline message states that the email address is already registered.",
                    "expected_result_plain": "An inline message states that the email address is already registered.",
                },
            ],
        }

    @staticmethod
    def graph_vault(root: Path, *, session_api="POST_save"):
        api_dir = root / "APIs" / "shop.test"
        param_dir = root / "Params"
        api_dir.mkdir(parents=True)
        param_dir.mkdir(parents=True)
        (api_dir / "POST_save.md").write_text(
            "## Params\n- [[Params/shop.test_checkout__email]]\n", encoding="utf-8"
        )
        (api_dir / "POST_other.md").write_text("## Captured API\n", encoding="utf-8")
        (param_dir / "shop.test_checkout__email.md").write_text(
            "## Used by\n- [[POST_save]] _request_\n", encoding="utf-8"
        )
        report_dir = root / "BRDs" / "reports"
        report_dir.mkdir(parents=True)
        (root / "BRDs" / "index").mkdir()
        (root / "BRDs" / "index" / "hub.md").write_text(
            "[[APIs/shop.test/POST_other]] [[BRDs/entries/unmatched]]", encoding="utf-8"
        )
        (report_dir / "links.json").write_text(json.dumps({
            "requirement_entity_links": [
                {"source_file": "BRD-test.md", "requirement_id": "BR-1", "entities": [
                    {"kind": "api", "path": "APIs/shop.test/POST_save.md", "match_basis": "exact_id"},
                ]},
                {"source_file": "BRD-test.md", "requirement_id": "BR-2", "entities": [
                    {"kind": "param", "path": "Params/shop.test_checkout__email.md", "match_basis": "field_name"},
                    {"kind": "api", "path": "APIs/shop.test/POST_save.md",
                     "match_basis": "api_used_by_matched_param"},
                ]},
                {"source_file": "BRD-test.md", "requirement_id": "BR-3", "entities": [
                    {"kind": "api", "path": "APIs/shop.test/missing.md", "match_basis": "exact_id"},
                ]},
                {"source_file": "BRD-test.md", "requirement_id": "BR-4", "entities": []},
            ]
        }), encoding="utf-8")
        connection = sqlite3.connect(":memory:")
        connection.execute("CREATE TABLE session_api_events (session_id TEXT, domain TEXT, ref_stem TEXT)")
        connection.execute(
            "INSERT INTO session_api_events VALUES (?, ?, ?)",
            ("session-a", "shop.test", session_api),
        )
        vault = SimpleNamespace(vault_path=root, conn=connection)
        contexts = [
            {"domain": "shop.test", "ref_stem": "POST_save", "family_ref_stems": ["POST_save"]},
            {"domain": "shop.test", "ref_stem": "POST_other", "family_ref_stems": ["POST_other"]},
        ]
        requirements = [
            {"source_file": "BRD-test.md", "requirement_id": f"BR-{index}",
             "traceability_key": f"BRD-test.md/BR-{index}"}
            for index in range(1, 5)
        ]
        return vault, contexts, requirements

    def test_graph_filter_includes_any_reachable_requirement_from_session_roots(self):
        with tempfile.TemporaryDirectory() as temp_dir:
            vault, contexts, requirements = self.graph_vault(Path(temp_dir))
            included, skipped = testcase_gen._filter_brd_requirements_by_graph(
                vault, contexts, requirements, session_id="session-a",
            )
            vault.conn.close()

        self.assertEqual([req["requirement_id"] for req in included], ["BR-1", "BR-2"])
        self.assertEqual(skipped, 2)

    def test_graph_filter_does_not_use_other_session_apis_or_navigation_hubs(self):
        with tempfile.TemporaryDirectory() as temp_dir:
            vault, contexts, requirements = self.graph_vault(Path(temp_dir), session_api="POST_other")
            included, skipped = testcase_gen._filter_brd_requirements_by_graph(
                vault, contexts, requirements, session_id="session-a",
            )
            vault.conn.close()

        self.assertEqual(included, [])
        self.assertEqual(skipped, 4)

    def test_disconnected_brd_requirements_are_not_sent_to_the_llm(self):
        req = self.requirement()
        with tempfile.TemporaryDirectory() as temp_dir:
            root = Path(temp_dir)
            connection = sqlite3.connect(":memory:")
            connection.execute("CREATE TABLE session_api_events (session_id TEXT, domain TEXT, ref_stem TEXT)")
            vault = SimpleNamespace(vault_path=root, conn=connection)
            with patch.object(testcase_gen, "build_contexts", return_value=[]), \
                    patch.object(testcase_gen, "_load_brd_requirements", return_value=[req]), \
                    patch.object(testcase_gen, "chat_json") as chat:
                rows = testcase_gen.generate_all(vault, brd_db_path="unused.sqlite")
            connection.close()

        self.assertEqual(rows, [])
        chat.assert_not_called()

    def test_debug_reports_graph_connected_requirement_counts(self):
        with tempfile.TemporaryDirectory() as temp_dir:
            root = Path(temp_dir)
            vault, _contexts, requirements = self.graph_vault(root)
            output = StringIO()
            with patch.object(testcase_gen, "build_contexts", return_value=[]), \
                    patch.object(testcase_gen, "_load_brd_requirements", return_value=requirements), \
                    redirect_stdout(output):
                testcase_gen.debug_dump(vault, brd_db_path="unused.sqlite")
            vault.conn.close()

        self.assertIn("BRD graph connectivity: 0/4 linked requirement(s)", output.getvalue())
        self.assertIn("4 disconnected or missing graph links skipped", output.getvalue())

    def test_requirement_cases_have_concrete_text_and_separate_composite_traceability(self):
        req = self.requirement()
        item = self.generated_item()
        with patch.object(testcase_gen, "chat_json", return_value={"items": [item]}) as chat:
            rows = testcase_gen._generate_brd_rows([], [req])

        self.assertEqual(len(rows), 2)
        self.assertEqual({row["Category"] for row in rows}, {"brd-positive", "brd-negative"})
        for row in rows:
            self.assertEqual(row["Traceability Key"], "BRD-01_Registration.md/BR-01.3")
            self.assertEqual(row["Applies to"], row["Traceability Key"])
            self.assertEqual(row["Priority"], "Must")
            self.assertEqual(row["Cycle"], "UAT, Regression, Integration")
            self.assertNotIn("BR-01.3", row["Test Case Name"])
            self.assertNotIn("entry-", row["Data"])
            self.assertRegex(row["Steps"], r"^1\. ")
            self.assertNotIn("requirement", row["Expected Result"].casefold())
        sent = json.loads(chat.call_args.args[1])
        self.assertEqual(sent["requirements"][0]["traceability_key"], "BRD-01_Registration.md/BR-01.3")

    def test_brd_batch_retries_only_missing_requirement(self):
        first = self.requirement()
        second = self.requirement(
            req_id="BR-02.1", source_file="BRD-02_Registration.md",
            statement="A customer can view their past orders.",
        )
        second_key = "BRD-02_Registration.md/BR-02.1"
        with patch.object(testcase_gen, "chat_json", side_effect=[
            {"items": [self.generated_item()]},
            {"items": [self.generated_item(second_key)]},
        ]) as chat:
            rows = testcase_gen._generate_brd_rows([], [first, second])

        self.assertEqual(chat.call_count, 2)
        self.assertEqual(len(rows), 4)
        retry_payload = json.loads(chat.call_args.args[1])
        self.assertEqual(
            [item["traceability_key"] for item in retry_payload["requirements"]],
            [second_key],
        )

    def test_exact_duplicates_merge_and_non_testable_entries_are_filtered(self):
        first = self.requirement()
        second = self.requirement(
            req_id="BR-1.3", source_file="BRD-01_Account_Security.md", priority="Should"
        )
        assumption = self.requirement(
            req_id="BRD-ASSUMPTION-1", entry_type="assumption",
            statement="Stakeholders are available for weekly review.",
        )
        deduplicated, conflicts, eligible = testcase_gen._prepare_brd_requirements(
            [first, second, assumption]
        )

        self.assertEqual(len(eligible), 2)
        self.assertEqual(len(deduplicated), 1)
        self.assertEqual(len(conflicts), 0)
        self.assertEqual(set(deduplicated[0]["applies_to"]), {
            "BRD-01_Registration.md/BR-01.3",
            "BRD-01_Account_Security.md/BR-1.3",
        })
        self.assertEqual(deduplicated[0]["priority"], "Must")
        self.assertFalse(deduplicated[0]["is_nfr"])

    def test_loader_keeps_composite_source_and_reads_distilled_entry_type(self):
        with tempfile.TemporaryDirectory() as temp_dir:
            root = Path(temp_dir)
            source = root / "BRD-01_User_Registration.md"
            source.touch()
            note = root / "requirement.md"
            note.write_text(
                "## Open questions\n- Is signup available to guests?\n\n## Source evidence\ntext\n",
                encoding="utf-8",
            )
            db_path = root / "ledger.sqlite"
            connection = sqlite3.connect(db_path)
            connection.executescript("""
                CREATE TABLE entries (
                    id TEXT, source_path TEXT, entry_type TEXT, title TEXT, content_json TEXT,
                    source_anchor_json TEXT, status TEXT, version TEXT, superseded_by TEXT
                );
                CREATE TABLE inventory (source_path TEXT, status TEXT);
            """)
            connection.execute(
                "INSERT INTO entries VALUES (?, ?, ?, ?, ?, ?, ?, ?, NULL)",
                (
                    "entry-1", str(source), "requirement", "BRD-ASSUMPTION-1",
                    json.dumps({
                        "type": "assumption", "id": "BRD-ASSUMPTION-1",
                        "requirement": "Assume weekly stakeholder review.",
                        "priority": "Unspecified", "markdown_path": str(note),
                    }),
                    json.dumps({"section": "Assumptions", "page": 3}), "unverified", "1.0",
                ),
            )
            connection.commit()
            connection.close()

            loaded = testcase_gen._load_brd_requirements(db_path)

        self.assertEqual(loaded[0]["entry_type"], "assumption")
        self.assertEqual(loaded[0]["source_file"], "BRD-01_User_Registration.md")
        self.assertEqual(
            loaded[0]["traceability_key"],
            "BRD-01_User_Registration.md/BRD-ASSUMPTION-1",
        )
        self.assertEqual(loaded[0]["open_questions"], ["Is signup available to guests?"])

    def test_missing_boundary_coverage_marks_requirement_blocked(self):
        req = self.requirement(
            statement="The password must contain at least 8 characters."
        )
        with patch.object(testcase_gen, "chat_json", return_value={
            "items": [self.generated_item()]
        }):
            rows = testcase_gen._generate_brd_rows([], [req])

        self.assertEqual(len(rows), 2)
        self.assertTrue(all(row["Review Status"] == "Blocked – needs clarification" for row in rows))
        self.assertTrue(all("boundary" in row["Review Note"] for row in rows))

    def test_colliding_ids_with_different_meanings_are_reported_not_generated(self):
        first = self.requirement(statement="The account email address is unique.")
        second = self.requirement(
            statement="The account password must be at least 8 characters.",
            source_file="BRD-01_Account_Security.md",
            req_id="BR-1.3",
        )
        rows = testcase_gen._generate_brd_rows([], [first, second])

        self.assertEqual(len(rows), 2)
        self.assertTrue(all(row["Output Sheet"] == "BRD Conflicts" for row in rows))
        self.assertEqual({row["BRD Requirement ID"] for row in rows}, {"BR-01.3", "BR-1.3"})

    def test_nfr_is_routed_to_measurable_separate_sheet(self):
        req = self.requirement(
            req_id="NFR-01", statement="The service shall be available 99.5% of each month.",
            source_file="BRD-10_Non_Functional.md",
        )
        item = {
            "traceability_key": "BRD-10_Non_Functional.md/NFR-01",
            "classification": "nfr",
            "review_status": "Ready for review",
            "nfr_metric": "Monthly service availability",
            "nfr_pass_criteria": "Availability is at least 99.5% for the calendar month.",
            "nfr_method_tool": "Calculate from production uptime monitoring and incident records.",
        }
        with patch.object(testcase_gen, "chat_json", return_value={"items": [item]}):
            rows = testcase_gen._generate_brd_rows([], [req])

        self.assertEqual(len(rows), 1)
        self.assertEqual(rows[0]["Output Sheet"], "NFR")
        self.assertEqual(rows[0]["Pass Criteria"], item["nfr_pass_criteria"])

    def test_unverified_or_open_question_case_is_blocked_in_review_metadata(self):
        req = self.requirement(status="unverified")
        req["open_questions"] = ["Should duplicate signup show inline text or a dialog?"]
        with patch.object(testcase_gen, "chat_json", return_value={
            "items": [self.generated_item()]
        }):
            rows = testcase_gen._generate_brd_rows([], [req])

        self.assertTrue(all(row["Review Status"] == "Blocked – needs clarification" for row in rows))
        self.assertTrue(all("Should duplicate signup" in row["Review Note"] for row in rows))
        self.assertTrue(all("Blocked" not in row["Expected Result"] for row in rows))

    def test_workbook_omits_non_testcase_records(self):
        rows = [
            {
                "Output Sheet": "NFR",
                "Traceability Key": "BRD-10.md/NFR-1",
                "Requirement": "Availability is 99.5%.",
                "Pass Criteria": "At least 99.5%.",
                "Method / Tool": "Uptime monitoring.",
            },
            {
                "Output Sheet": "BRD Conflicts",
                "Collision Key": "br13",
                "BRD Requirement ID": "BR-1.3",
                "BRD Evidence": "Different meaning.",
            },
            {
                "Output Sheet": "BRD Review",
                "Traceability Key": "BRD-01.md/BR-1.1",
                "Review Status": "Blocked – needs clarification",
            },
        ]
        with tempfile.TemporaryDirectory() as temp_dir:
            path = testcase_gen.write_excel(rows, Path(temp_dir) / "cases.xlsx")
            from openpyxl import load_workbook

            workbook = load_workbook(path, data_only=True)
            self.assertEqual(workbook.sheetnames, ["Test Cases"])
            self.assertEqual(workbook["Test Cases"].max_row, 5)
            workbook.close()

    def test_brd_cases_use_the_single_tester_facing_schema(self):
        req = self.requirement(status="verified")
        item = self.generated_item()
        item["module"] = "Accounts"
        with patch.object(testcase_gen, "chat_json", return_value={
            "items": [item]
        }):
            rows = testcase_gen._generate_brd_rows([], [req])
        with tempfile.TemporaryDirectory() as temp_dir:
            path = testcase_gen.write_excel(rows, Path(temp_dir) / "cases.xlsx")
            from openpyxl import load_workbook

            workbook = load_workbook(path, data_only=True)
            sheet = workbook["Test Cases"]
            headers = [sheet.cell(5, column).value for column in range(1, 12)]
            values = dict(zip(headers, next(sheet.iter_rows(min_row=6, max_row=6, values_only=True))))
            self.assertEqual(workbook.sheetnames, ["Test Cases"])
            for field in ("Test Case Name", "Precondition", "Data", "Steps", "Expected Result"):
                value = values[field] or ""
                self.assertNotIn("BR-01.3", value)
                self.assertNotIn("BRD-01_Registration.md", value)
                self.assertNotIn("**", value)
            self.assertEqual(values["Status"], "Untested")
            self.assertIn(values["Actual Result"], ("", None))
            self.assertIn(values["Evidence/Note"], ("", None))
            workbook.close()


class OpenRouterErrorTests(unittest.TestCase):
    def test_raw_output_is_available_without_expanding_error_message(self):
        output = "x" * 900
        error = OpenRouterError(f"bad JSON\nRaw output: {output[:500]}", raw_output=output)

        self.assertEqual(error.raw_output, output)
        self.assertLess(len(str(error)), len(output))


if __name__ == "__main__":
    unittest.main()
