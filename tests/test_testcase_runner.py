import hashlib
import json
import tempfile
import time
import unittest
from unittest.mock import patch
from pathlib import Path

from openpyxl import load_workbook

from capture_app.workflows.case_metadata import (
    decode_payload, encode_payload, load_metadata_sheet, row_hash, write_metadata_sheet,
)
from capture_app.workflows import testcase_gen
from capture_app.workflows.testcase_runner import (
    TestcaseRunner, _actual_result, _evaluate_structured, _normalize_number_or_date,
    _norm, _redact_observation, _retryable_infrastructure_error,
)


class FakeRequest:
    resource_type = "xhr"

    def __init__(self, url, method="POST"):
        self.url = url
        self.method = method
        self.failure = None

    def is_navigation_request(self):
        return False


class FakeResponse:
    def __init__(self, request, status):
        self.request = request
        self.url = request.url
        self.status = status


class FakeLocator:
    def __init__(self, page):
        self.page = page

    @property
    def first(self):
        return self

    def nth(self, _index):
        return self

    def click(self, **_kwargs):
        if self.page.on_click:
            self.page.on_click(self.page)

    def fill(self, _value, **_kwargs):
        pass

    def check(self, **_kwargs):
        pass

    def uncheck(self, **_kwargs):
        pass

    def select_option(self, **_kwargs):
        pass


class FakePage:
    def __init__(self, on_click=None):
        self.url = "about:blank"
        self.title = "Store"
        self.handlers = {}
        self.on_click = on_click
        self.clicks = 0
        self.observation = {
            "url": "https://shop.example.test/", "title": "Store", "headings": [],
            "messages": [], "field_states": [], "transient_messages": [],
        }

    def on(self, name, callback):
        self.handlers.setdefault(name, []).append(callback)

    def goto(self, url, **_kwargs):
        self.url = url
        self.observation["url"] = url

    def wait_for_timeout(self, _milliseconds):
        pass

    def wait_for_load_state(self, _state, **_kwargs):
        pass

    def get_by_role(self, _role, **_kwargs):
        return FakeLocator(self)

    def evaluate(self, _script):
        return json.loads(json.dumps(self.observation, ensure_ascii=False))

    def emit_outcome(self, message, status=200, url="https://shop.example.test/api/action"):
        request = FakeRequest(url)
        for callback in self.handlers.get("request", []):
            callback(request)
        self.observation["messages"] = [message]
        response = FakeResponse(request, status)
        for callback in self.handlers.get("response", []):
            callback(response)


class FakeRoute:
    def __init__(self, request):
        self.request = request
        self.aborted = False

    def abort(self):
        self.aborted = True

    def continue_(self):
        pass


class FakeContext:
    def __init__(self, on_click=None):
        self.page = FakePage(on_click)
        self.closed = False
        self.init_scripts = []
        self.routes = []

    def add_init_script(self, script):
        self.init_scripts.append(script)

    def route(self, pattern, handler):
        self.routes.append((pattern, handler))

    def new_page(self):
        return self.page

    def close(self):
        self.closed = True


class FakeBrowser:
    def __init__(self, on_click=None):
        self.on_click = on_click
        self.contexts = []

    def new_context(self, **_kwargs):
        context = FakeContext(self.on_click)
        self.contexts.append(context)
        return context


def claim(text, *, kind="message", quoted=None, status=None, url=None, field=None, value=None):
    return {"text": text, "kind": kind, "literals": {
        "quoted_message_text": quoted, "expected_status_class": status,
        "url_pattern": url, "field_name": field, "field_value": value,
    }}


def make_row(case_id, *, actions=None, claims=None, expected="The page shows a result.", prerequisites=None,
             cleanup=None):
    actions = actions or [{"kind": "click", "role": "button", "name": "View order", "value": ""}]
    return {
        "Module": "Orders", "Sub-module": "Order submit", "Test Case Name": case_id,
        "Precondition": "A test order is ready.", "Data": "Use isolated test data.",
        "Steps": "1. Click View order", "Expected Result": expected,
        "Category": "observed-flow", "_transition_key": case_id,
        "_claims": claims or [claim(expected, quoted="Order saved")],
        "_runner_payload": {
            "case_id": case_id, "source_session_id": "session-1", "source_run_id": "crawl-1",
            "start_url": "https://shop.example.test/", "route_actions": actions,
            "test_data": {}, "observed_checks": {}, "prerequisites": prerequisites or [],
            "cleanup": cleanup or [], "claims": claims or [claim(expected, quoted="Order saved")],
        },
    }


def write_cases(path: Path, rows):
    testcase_gen.write_excel(rows, path)
    return path


class CaseMetadataTests(unittest.TestCase):
    def test_payload_round_trips_after_chunking_and_row_hash_covers_expected_result(self):
        import random
        randomizer = random.Random(77)
        payload = {"data": "".join(randomizer.choice("abcdefghijklmnopqrstuvwxyz0123456789") for _ in range(85000))}
        encoded, digest = encode_payload(payload)
        self.assertEqual(decode_payload([encoded[i:i + 30000] for i in range(0, len(encoded), 30000)], digest), payload)
        before = row_hash({"Expected Result": "Order saved"}, payload)
        self.assertNotEqual(before, row_hash({"Expected Result": "Order failed"}, payload))

    def test_metadata_sheet_round_trips_happy_and_branch_payloads(self):
        from openpyxl import Workbook
        workbook = Workbook()
        cases = [
            {"case_id": "s:happy", "sheet_row": 6, "payload": {"route_actions": [{"name": "Open"}]},
             "row_sha256": "a"},
            {"case_id": "s:r:4", "sheet_row": 7, "payload": {"route_actions": [{"name": "Save"}]},
             "row_sha256": "b"},
        ]
        write_metadata_sheet(workbook, cases)
        loaded = load_metadata_sheet(workbook)
        self.assertEqual(loaded["s:happy"]["payload"], cases[0]["payload"])
        self.assertEqual(loaded["s:r:4"]["payload"], cases[1]["payload"])


class ClaimComparisonTests(unittest.TestCase):
    def test_navigation_uses_bounded_dom_content_wait(self):
        from capture_app.crawling.replayer import goto_once

        class RecordingPage:
            def __init__(self):
                self.calls = []

            def goto(self, url, **kwargs):
                self.calls.append((url, kwargs))

        page = RecordingPage()
        goto_once(page, "https://shop.example.test/")
        self.assertEqual(page.calls, [("https://shop.example.test/", {
            "wait_until": "domcontentloaded", "timeout": 15000,
        })])

    def test_vietnamese_normalization_preserves_display_text_but_compares_nfc(self):
        composed = chr(0x0110) + chr(0x00e3) + " l" + chr(0x01b0) + "u h" + chr(0x1ed3) + " s" + chr(0x01a1) + " th" + chr(0x00e0) + "nh c" + chr(0x00f4) + "ng"
        decomposed = chr(0x0110) + "a\u0303 lu\u031bu ho\u0302\u0300 so\u031b tha\u0300nh co\u0302ng"
        self.assertEqual(_norm(composed), _norm(decomposed))
        rendered = _actual_result({"action_messages": [composed], "title": "Store", "url": "https://shop.example.test", "own_request": {"responses": []}})
        self.assertIn(composed, rendered)
        self.assertIn(chr(0x0110) + chr(0x00e3), rendered)

    def test_invalid_url_actual_result_explains_failure_in_plain_language(self):
        rendered = _actual_result(
            {"url": "about:blank", "title": "", "own_request": {"responses": []}},
            "Error: Page.goto: Cannot navigate to invalid URL",
        )
        self.assertIn("could not open the saved website address", rendered)
        self.assertIn("blank page", rendered)
        self.assertNotIn("Messages seen:", rendered)

    def test_redaction_hides_fixture_and_email_without_changing_other_casing(self):
        observation = {"url": "https://shop.example.test", "title": "Store", "messages": ["Order Alice@example.test Saved"],
                       "headings": [], "field_states": [], "transient_messages": [], "own_request": {}}
        redacted = _redact_observation(observation, ("fixture-secret",))
        rendered = _actual_result(redacted)
        self.assertIn("Order", rendered)
        self.assertIn("Saved", rendered)
        self.assertNotIn("Alice@example.test", rendered)
        self.assertNotIn("fixture-secret", rendered)

    def test_vietnamese_numbers_and_dates_normalize_equivalently(self):
        self.assertEqual(_normalize_number_or_date("1.234,50"), _normalize_number_or_date("1,234.50"))
        self.assertEqual(_normalize_number_or_date("1.234"), _normalize_number_or_date("1234"))
        self.assertEqual(_normalize_number_or_date("31/10/2026"), "2026-10-31")

    def test_deterministic_missing_message_is_not_observed_and_4xx_vetoes_success(self):
        runner = TestcaseRunner("placeholder.xlsx", semantic="off")
        obs = {"messages": [], "action_messages": [], "field_states": [], "headings": [], "url": "https://shop.example.test/",
               "own_request": {"requests": [], "responses": [], "failures": []}}
        status, _, _ = runner._evaluate_case("Order saved", [claim("Order saved", quoted="Order saved")], obs, "a")
        self.assertEqual(status, "Unknown")
        obs["own_request"]["responses"] = [{"method": "POST", "url": "https://shop.example.test/api", "status": 422}]
        status, results, _ = runner._evaluate_case("Order saved successfully", [claim("Order saved successfully", kind="other")], obs, "a")
        self.assertEqual(status, "Failed")
        self.assertEqual(results[0]["basis"], "structured")

    def test_semantic_modes_and_verbatim_evidence_guard(self):
        observation = {"messages": ["Order saved"], "action_messages": ["Order saved"], "title": "Store", "headings": [],
                       "field_states": [], "url": "https://shop.example.test/",
                       "own_request": {"requests": [], "responses": [], "failures": []}}
        calls = []
        judge = lambda _claim, _obs: calls.append(1) or {"verdict": "supported", "evidence_span": "Order saved"}
        one_claim = [claim("Giao dịch hoàn tất", kind="message")]
        off = TestcaseRunner("placeholder.xlsx", semantic="off", judge=judge)
        self.assertEqual(off._evaluate_case("Order saved", one_claim, observation, "a")[0], "Unknown")
        self.assertEqual(calls, [])
        advisory = TestcaseRunner("placeholder.xlsx", semantic="advisory", judge=judge)
        status, results, note = advisory._evaluate_case("Order saved", one_claim, observation, "a")
        self.assertEqual(status, "Unknown")
        self.assertIn("SEMANTIC-ONLY: suggested=Passed", note)
        self.assertEqual(results[0]["verdict"], "supported")
        on = TestcaseRunner("placeholder.xlsx", semantic="on", judge=judge)
        self.assertEqual(on._evaluate_case("Giao dịch hoàn tất", one_claim, observation, "a")[0], "Passed")
        bad_judge = TestcaseRunner("placeholder.xlsx", semantic="on", judge=lambda *_: {
            "verdict": "supported", "evidence_span": "not in observation"})
        result = bad_judge._semantic_judge(one_claim[0], observation)
        self.assertEqual(result["verdict"], "not_observed")
        self.assertGreater(len(calls), 0)

    def test_claim_literal_run_id_is_substituted_before_check(self):
        runner = TestcaseRunner("placeholder.xlsx", semantic="off")
        runner.run_id = "run-17"
        observation = {"messages": [], "action_messages": ["Order run-17 saved"], "title": "Store", "headings": [],
                       "field_states": [], "url": "https://shop.example.test/",
                       "own_request": {"requests": [], "responses": [], "failures": []}}
        status, results, _ = runner._evaluate_case(
            "Order {{run_id}} saved", [claim("Order {{run_id}} saved", quoted="Order {{run_id}} saved")], observation, "a")
        self.assertEqual(status, "Passed")
        self.assertEqual(results[0]["verdict"], "supported")

    def test_opposite_polarity_and_wrong_entity_or_number_do_not_pass(self):
        observed = {"messages": ["Saved order 417 for Alice"], "action_messages": ["Saved order 417 for Alice"],
                    "field_states": [], "headings": [], "url": "https://shop.example.test/",
                    "own_request": {"requests": [], "responses": [], "failures": []}}
        runner = TestcaseRunner("placeholder.xlsx", semantic="off")
        status, _, _ = runner._evaluate_case(
            "Order 471 for Alice failed", [claim("Order 471 for Alice failed", quoted="Order 471 for Alice failed")], observed, "a")
        self.assertNotEqual(status, "Passed")
        observed["messages"] = ["Not saved"]
        observed["action_messages"] = ["Not saved"]
        status, _, _ = runner._evaluate_case(
            "Saved", [claim("Saved", quoted="Saved")], observed, "a")
        self.assertEqual(status, "Failed")
        status, _, _ = runner._evaluate_case(
            "Order 417 for Bob saved", [claim("Order 417 for Bob saved", quoted="Order 417 for Bob saved")], observed, "a")
        self.assertNotEqual(status, "Passed")

    def test_calibrated_semantic_opposite_polarity_is_failed_only_with_stubbed_judge(self):
        observation = {"messages": ["Save failed"], "action_messages": ["Save failed"], "field_states": [],
                       "headings": [], "url": "https://shop.example.test/",
                       "own_request": {"requests": [], "responses": [], "failures": []}}
        runner = TestcaseRunner("placeholder.xlsx", semantic="on", judge=lambda *_: {
            "verdict": "contradicted", "evidence_span": "Save failed"})
        status, _, _ = runner._evaluate_case("Save succeeded", [claim("Save succeeded", kind="message")], observation, "a")
        self.assertEqual(status, "Failed")

    def test_identical_evidence_reuses_cached_judge_answer(self):
        class StubResponseCache:
            def __init__(self):
                self.entries = {}
                self.calls = 0

            def chat_json(self, _system, prompt, **_kwargs):
                if prompt not in self.entries:
                    self.calls += 1
                    self.entries[prompt] = {"verdict": "supported", "evidence_span": "Order saved"}
                return self.entries[prompt]

        runner = TestcaseRunner("placeholder.xlsx", semantic="on")
        cache = StubResponseCache()
        runner._response_cache = cache
        observation = {"messages": ["Order saved"], "action_messages": ["Order saved"], "field_states": [],
                       "headings": [], "url": "https://shop.example.test/",
                       "own_request": {"requests": [], "responses": [], "failures": []}}
        one_claim = [claim("Vietnamese wording", kind="message")]
        first = runner._evaluate_case("Vietnamese wording", one_claim, observation, "a")[0]
        second = runner._evaluate_case("Vietnamese wording", one_claim, observation, "a")[0]
        self.assertEqual((first, second), ("Passed", "Passed"))
        self.assertEqual(cache.calls, 1)

    def test_claim_literals_can_use_fixture_test_data(self):
        runner = TestcaseRunner("placeholder.xlsx", semantic="off", fixtures={"ORDER_CODE": "ORD-2026-17"})
        runner.run_id = "run-17"
        observation = {"messages": ["Created ORD-2026-17"], "action_messages": ["Created ORD-2026-17"],
                       "field_states": [], "headings": [], "url": "https://shop.example.test/",
                       "own_request": {"requests": [], "responses": [], "failures": []}}
        status, _, _ = runner._evaluate_case("Created {{ORDER_CODE}}", [claim(
            "Created {{ORDER_CODE}}", quoted="Created {{ORDER_CODE}}")], observation, "attempt-1")
        self.assertEqual(status, "Passed")

    def test_unrelated_toast_does_not_supply_outcome_type_for_other_claim(self):
        runner = TestcaseRunner("placeholder.xlsx", semantic="off")
        observation = {"messages": ["Unrelated saved toast"], "action_messages": ["Unrelated saved toast"],
                       "title": "Store", "headings": [], "field_states": [], "url": "https://shop.example.test/",
                       "own_request": {"requests": [], "responses": [], "failures": []}}
        status, _, _ = runner._evaluate_case("The record is complete", [claim("The record is complete", kind="other")], observation, "a")
        self.assertEqual(status, "Unknown")

    def test_supported_url_claim_alone_is_unknown_without_no_oracle_code(self):
        runner = TestcaseRunner("placeholder.xlsx", semantic="off")
        observation = {"messages": [], "action_messages": [], "field_states": [], "headings": [],
                       "url": "https://shop.example.test/orders/42", "own_request": {"responses": []}}
        status, results, note = runner._evaluate_case("The order page is shown", [claim(
            "The order page is shown", kind="url", url="/orders/42")], observation, "a")
        self.assertEqual(status, "Unknown")
        self.assertEqual(results[0]["verdict"], "supported")
        self.assertEqual(note, "")


class RunnerWorkbookTests(unittest.TestCase):
    def setUp(self):
        self.tmp = tempfile.TemporaryDirectory()
        self.addCleanup(self.tmp.cleanup)
        self.root = Path(self.tmp.name)

    @staticmethod
    def successful_click(page):
        page.emit_outcome("Order saved")

    def test_runner_uses_fresh_contexts_and_writes_results_copy_only(self):
        source = self.root / "cases.xlsx"
        rows = [make_row("case-1"), make_row("case-2")]
        write_cases(source, rows)
        before = hashlib.sha256(source.read_bytes()).hexdigest()
        browser = FakeBrowser(self.successful_click)
        contexts = []

        def factory(_auth_state):
            context = browser.new_context()
            contexts.append(context)
            return browser, context.close, context

        runner = TestcaseRunner(source, browser_factory=factory, semantic="off")
        result = runner.run()
        self.assertEqual(len(contexts), 2)
        self.assertTrue(all(context.closed for context in contexts))
        self.assertEqual(hashlib.sha256(source.read_bytes()).hexdigest(), before)
        result_book = load_workbook(result["output_path"])
        self.assertEqual(result_book.sheetnames, ["Test Cases", "Runner Metadata"])
        self.assertEqual(result_book["Runner Metadata"].sheet_state, "hidden")
        result_sheet = result_book["Test Cases"]
        self.assertEqual(result_sheet.column_dimensions["F"].width, 72)
        self.assertTrue(result_sheet.data_validations.dataValidation)
        source_book = load_workbook(source)
        self.assertEqual(result_sheet["A5"].fill.fgColor.rgb, source_book["Test Cases"]["A5"].fill.fgColor.rgb)
        source_book.close()
        self.assertEqual(result_sheet["I6"].value,
                         "The page showed “Order saved”.\n"
                         "The browser reached “Store” at https://shop.example.test/.\n"
                         "The action received HTTP 200 from POST https://shop.example.test/api/action.")
        self.assertEqual(result_sheet["J6"].value, "Passed")
        result_book.close()

    def test_rerun_failed_reuses_run_id_and_does_not_replay_passed_rows(self):
        source = self.root / "rerun.xlsx"
        rows = [make_row("case-pass", claims=[claim("Order saved", quoted="Order saved")]),
                make_row("case-fail", claims=[claim("Order saved", quoted="Order saved")])]
        write_cases(source, rows)
        clicks = []

        def outcome(page):
            clicks.append(1)
            page.emit_outcome("Order saved" if len(clicks) in {1, 3} else "Save failed")

        first = TestcaseRunner(source, browser_factory=lambda _auth: self._fake_browser(outcome), semantic="off").run()
        records = [json.loads(line) for line in Path(first["report_path"]).read_text(encoding="utf-8").splitlines()]
        initial_run_id = records[0]["run_id"]
        self.assertEqual(len(clicks), 2)
        second = TestcaseRunner(source, browser_factory=lambda _auth: self._fake_browser(outcome),
                                semantic="off", rerun_failed=True).run()
        records = [json.loads(line) for line in Path(second["report_path"]).read_text(encoding="utf-8").splitlines()]
        self.assertEqual(records[0]["run_id"], initial_run_id)
        self.assertEqual(len(clicks), 3)

    def test_rerun_failed_also_selects_replay_error_rows(self):
        source = self.root / "rerun-replay-error.xlsx"
        write_cases(source, [make_row("replay-error", claims=[claim("Order saved", quoted="Order saved")])])
        calls = []

        def first_error_then_success(page):
            calls.append(1)
            if len(calls) <= 2:
                raise ValueError("unsupported selector")
            page.emit_outcome("Order saved")

        first = TestcaseRunner(source, browser_factory=lambda _auth: self._fake_browser(first_error_then_success),
                               semantic="off").run()
        first_case = next(json.loads(line) for line in Path(first["report_path"]).read_text(encoding="utf-8").splitlines()
                          if json.loads(line).get("type") == "case")
        self.assertEqual(first_case["outcome_code"], "REPLAY-ERROR:")
        second = TestcaseRunner(source, browser_factory=lambda _auth: self._fake_browser(first_error_then_success),
                                semantic="off", rerun_failed=True).run()
        self.assertEqual(len(calls), 3)
        records = [json.loads(line) for line in Path(second["report_path"]).read_text(encoding="utf-8").splitlines()]
        latest_case = next(record for record in reversed(records) if record.get("type") == "case")
        self.assertEqual(latest_case["status"], "Passed")

    def test_safety_preflight_holds_intermediate_submit_without_opening_browser(self):
        source = self.root / "held.xlsx"
        actions = [
            {"kind": "click", "role": "button", "name": "Continue", "value": ""},
            {"kind": "click", "role": "button", "name": "Submit order", "value": ""},
        ]
        write_cases(source, [make_row("held", actions=actions)])
        launched = []
        runner = TestcaseRunner(source, browser_factory=lambda *_: launched.append(1), semantic="off")
        result = runner.run()
        book = load_workbook(result["output_path"])
        sheet = book["Test Cases"]
        self.assertEqual(sheet["J6"].value, "Untested")
        self.assertIn(sheet["I6"].value, ("", None))
        self.assertIn("HELD:", sheet["K6"].value)
        self.assertEqual(launched, [])
        book.close()

    def test_hash_mismatch_marks_edited_and_does_not_execute(self):
        source = self.root / "edited.xlsx"
        write_cases(source, [make_row("edited")])
        book = load_workbook(source)
        book["Test Cases"]["G6"] = "Changed expected result"
        book.save(source)
        book.close()
        launched = []
        result = TestcaseRunner(source, browser_factory=lambda *_: launched.append(1), semantic="off").run()
        output = load_workbook(result["output_path"])
        self.assertEqual(output["Test Cases"]["J6"].value, "Untested")
        self.assertIn(output["Test Cases"]["I6"].value, ("", None))
        self.assertIn("EDITED:", output["Test Cases"]["K6"].value)
        self.assertEqual(launched, [])
        output.close()

    def test_allow_edited_replays_metadata_payload_and_flags_evidence(self):
        source = self.root / "edited-allowed.xlsx"
        write_cases(source, [make_row("edited-allowed")])
        book = load_workbook(source)
        book["Test Cases"]["G6"] = "Edited visible expected text"
        book.save(source)
        book.close()
        launched = []
        result = TestcaseRunner(source, browser_factory=lambda _auth: (launched.append(1) or self._fake_browser(self.successful_click)),
                                allow_edited=True, semantic="off").run()
        output = load_workbook(result["output_path"])
        self.assertTrue(launched)
        self.assertIn("EDITED: replay used stored metadata payload", output["Test Cases"]["K6"].value)
        output.close()

    def test_default_allowed_origins_are_derived_from_route_and_checks(self):
        runner = TestcaseRunner("placeholder.xlsx", semantic="off")
        payload = {"start_url": "https://shop.example.test/", "route_actions": [
            {"kind": "goto", "name": "https://idp.example.test/login"}],
            "observed_checks": {"after_url": "https://shop.example.test/done"}}
        self.assertEqual(runner._allowed_origins(payload), {"https://shop.example.test"})
        runner.allowed_hosts = ["idp.example.test"]
        self.assertEqual(runner._allowed_origins(payload), {
            "https://shop.example.test", "https://idp.example.test",
        })

    def test_legacy_host_only_start_url_is_normalized_for_replay_and_host_guard(self):
        runner = TestcaseRunner("placeholder.xlsx", semantic="off")
        payload = {"start_url": "demowebshop.tricentis.com/", "route_actions": [
            {"kind": "click", "role": "button", "name": "Log in", "value": ""},
        ]}
        actions, issue = runner._preflight(payload, "attempt-1")
        self.assertEqual(issue, "")
        self.assertEqual(runner._prepared_start_url, "https://demowebshop.tricentis.com/")
        self.assertEqual(runner._allowed_origins(payload), {"https://demowebshop.tricentis.com"})
        self.assertEqual(len(actions), 1)

    def test_infers_captured_source_page_when_route_has_no_navigation_step(self):
        runner = TestcaseRunner("placeholder.xlsx", semantic="off")
        payload = {"start_url": "https://shop.example.test/", "observed_checks": {
            "before_url": "https://shop.example.test/blue-jeans",
        }, "route_actions": [
            {"kind": "fill", "role": "textbox", "name": "", "value": "test"},
        ]}
        actions, issue = runner._preflight(payload, "attempt-1")
        self.assertEqual(issue, "")
        self.assertEqual(runner._prepared_start_url, "https://shop.example.test/blue-jeans")
        self.assertEqual(len(actions), 1)

    def test_replay_starts_at_captured_source_page_before_performing_action(self):
        source = self.root / "source-page-entry.xlsx"
        row = make_row("source-page", actions=[
            {"kind": "click", "role": "button", "name": "Open", "value": ""},
        ])
        row["_runner_payload"]["observed_checks"] = {
            "before_url": "https://shop.example.test/blue-jeans",
        }
        write_cases(source, [row])
        clicked_urls = []

        def record_page(page):
            clicked_urls.append(page.url)
            self.successful_click(page)

        TestcaseRunner(source, browser_factory=lambda _auth: self._fake_browser(record_page),
                       semantic="off").run()
        self.assertEqual(clicked_urls, ["https://shop.example.test/blue-jeans"])

    def test_keeps_start_page_when_captured_click_path_reaches_source(self):
        runner = TestcaseRunner("placeholder.xlsx", semantic="off")
        payload = {"start_url": "https://shop.example.test/", "observed_checks": {
            "before_url": "https://shop.example.test/newproducts",
        }, "route_actions": [
            {"kind": "click", "role": "link", "name": "New products", "value": ""},
            {"kind": "click", "role": "link", "name": "Phone Cover", "value": ""},
        ]}
        actions, issue = runner._preflight(payload, "attempt-1")
        self.assertEqual(issue, "")
        self.assertEqual(runner._prepared_start_url, "https://shop.example.test/")
        self.assertEqual(len(actions), 2)

    def test_prerequisite_uses_provider_action_request_not_judged_status(self):
        runner = TestcaseRunner("placeholder.xlsx")
        payload = {"prerequisites": [{"case_id": "provider"}]}
        provider = {"case_id": "provider", "status": "Unknown", "action_request_succeeded": True}
        self.assertEqual(runner._case_prerequisite_issue(payload, {"provider": provider}, "dependent"), "")
        self.assertIn("PREREQ-UNMET", runner._case_prerequisite_issue(payload, {}, "dependent"))
        self.assertEqual(runner._case_prerequisite_issue(
            {"prerequisites": [{"case_id": "provider", "fixture": "ORDER_ID"}]}, {}, "dependent"),
            "PREREQ-UNMET: provider provider has no successful action request in this run or fixture")
        runner.fixtures["ORDER_ID"] = "123"
        self.assertEqual(runner._case_prerequisite_issue(
            {"prerequisites": [{"case_id": "provider", "fixture": "ORDER_ID"}]}, {}, "dependent"), "")

    def test_tracking_beacon_is_excluded_from_action_owned_requests(self):
        runner = TestcaseRunner("placeholder.xlsx")
        runner._active_request_capture = True
        runner._action_origin = "https://shop.example.test"
        runner._action_requests = []
        runner._action_responses = []
        runner._on_request(FakeRequest("https://tracker.example.test/pixel"), FakePage())
        runner._on_response(FakeResponse(FakeRequest("https://tracker.example.test/pixel"), 500), FakePage())
        self.assertEqual(runner._action_requests, [])
        self.assertEqual(runner._action_responses, [])
        observation = {"messages": ["Order saved"], "action_messages": ["Order saved"], "field_states": [],
                       "headings": [], "url": "https://shop.example.test/",
                       "own_request": {"requests": [], "responses": runner._action_responses, "failures": []}}
        self.assertEqual(runner._evaluate_case(
            "Order saved", [claim("Order saved", quoted="Order saved")], observation, "a")[0], "Passed")

    def test_same_origin_action_request_4xx_veto_ignores_third_party_beacon(self):
        runner = TestcaseRunner("placeholder.xlsx", semantic="off")
        observation = {"messages": ["Order saved"], "action_messages": ["Order saved"], "field_states": [],
                       "headings": [], "url": "https://shop.example.test/",
                       "own_request": {"requests": [], "failures": [], "responses": [
                           {"method": "POST", "url": "https://shop.example.test/api/orders", "status": 201},
                       ]}}
        status, _, _ = runner._evaluate_case("Order saved", [claim("Order saved", quoted="Order saved")], observation, "a")
        self.assertEqual(status, "Passed")

    def test_4xx_on_mutating_request_vetoes_even_when_other_read_succeeds(self):
        runner = TestcaseRunner("placeholder.xlsx", semantic="off")
        observation = {"messages": ["Order saved"], "action_messages": ["Order saved"], "field_states": [],
                       "headings": [], "url": "https://shop.example.test/",
                       "own_request": {"requests": [], "failures": [], "responses": [
                           {"method": "GET", "url": "https://shop.example.test/api/sidebar", "status": 200},
                           {"method": "POST", "url": "https://shop.example.test/api/orders", "status": 422},
                       ]}}
        status, _, _ = runner._evaluate_case("Order saved", [claim("Order saved", quoted="Order saved")], observation, "a")
        self.assertEqual(status, "Failed")

    def test_default_navigation_route_aborts_external_redirect(self):
        runner = TestcaseRunner("placeholder.xlsx", semantic="off")
        payload = {"start_url": "https://shop.example.test/", "route_actions": [
            {"kind": "click", "role": "button", "name": "View", "value": ""}]}
        browser = FakeBrowser()
        context = FakeContext()
        runner.browser_factory = lambda _auth: (browser, context.close, context)
        actions, issue = runner._preflight(payload, "attempt-1")
        self.assertEqual(issue, "")
        runner.run_id = "run-1"
        runner._capture_case(payload, actions, "attempt-1")
        handler = context.routes[0][1]
        route = FakeRoute(FakeRequest("https://evil.example.test/redirect"))
        route.request.is_navigation_request = lambda: True
        handler(route)
        self.assertTrue(route.aborted)

    def test_attempt_id_template_is_shared_by_every_route_step(self):
        runner = TestcaseRunner("placeholder.xlsx", semantic="off")
        runner.run_id = "run-1"
        payload = {"start_url": "https://shop.example.test/", "route_actions": [
            {"kind": "fill", "role": "textbox", "name": "Order", "value": "{{attempt_id}}"},
            {"kind": "click", "role": "button", "name": "View", "value": ""},
        ]}
        actions, issue = runner._preflight(payload, "attempt-17")
        self.assertEqual(issue, "")
        self.assertEqual(actions[0].value, "attempt-17")

    def test_retry_failed_selection_includes_replay_error_only_for_unknown(self):
        self.assertTrue(_retryable_infrastructure_error("Timeout while waiting for navigation"))
        self.assertTrue(_retryable_infrastructure_error("Execution context was destroyed"))
        self.assertFalse(_retryable_infrastructure_error("claim contradicted"))

    def test_no_metadata_rows_are_marked_no_fixture_and_actual_result_stays_blank(self):
        source = self.root / "legacy.xlsx"
        write_cases(source, [make_row("legacy")])
        book = load_workbook(source)
        del book["Runner Metadata"]
        book.save(source)
        book.close()
        result = TestcaseRunner(source, semantic="off").run()
        self.assertEqual(result["cases_replayed"], 0)
        output = load_workbook(result["output_path"])
        self.assertEqual(output["Test Cases"]["J6"].value, "Untested")
        self.assertIn("NO-FIXTURE:", output["Test Cases"]["K6"].value)
        self.assertIn(output["Test Cases"]["I6"].value, (None, ""))
        output.close()

    def test_browser_start_failure_is_reported_once_without_browser_none_cascade(self):
        source = self.root / "browser-start-failure.xlsx"
        write_cases(source, [make_row("case-1"), make_row("case-2")])

        class FakeChromium:
            launches = 0

            def launch(self, **_kwargs):
                self.launches += 1
                raise RuntimeError("Executable doesn't exist at chromium.exe")

        class FakePlaywright:
            def __init__(self):
                self.chromium = FakeChromium()
                self.stopped = False

            def stop(self):
                self.stopped = True

        fake_playwright = FakePlaywright()
        with patch("playwright.sync_api.sync_playwright") as sync_playwright:
            sync_playwright.return_value.start.return_value = fake_playwright
            result = TestcaseRunner(source, semantic="off").run()

        self.assertEqual(fake_playwright.chromium.launches, 1)
        self.assertTrue(fake_playwright.stopped)
        self.assertEqual(result["cases_replayed"], 0)
        self.assertIn("Executable doesn't exist", result["browser_start_error"])
        output = load_workbook(result["output_path"])
        sheet = output["Test Cases"]
        for row_number in (6, 7):
            self.assertEqual(sheet.cell(row_number, 10).value, "Unknown")
            self.assertIn("Browser startup failed", sheet.cell(row_number, 11).value)
            self.assertNotIn("AttributeError", sheet.cell(row_number, 11).value)
        output.close()

    def test_transient_toast_is_written_even_after_it_disappears(self):
        source = self.root / "toast.xlsx"
        write_cases(source, [make_row("toast")])

        def transient(page):
            page.observation["messages"] = []
            page.observation["transient_messages"] = [{"text": "Order saved", "ts": int(time.time() * 1000),
                                                        "role": "status", "ariaLive": "polite", "className": "toast"}]

        contexts = []

        def factory(_auth):
            created = self._fake_browser(transient)
            contexts.append(created[2])
            return created

        result = TestcaseRunner(source, browser_factory=factory, semantic="off").run()
        output = load_workbook(result["output_path"])
        sheet = output["Test Cases"]
        self.assertEqual(sheet["J6"].value, "Passed")
        self.assertIn("Order saved", sheet["I6"].value)
        self.assertIn("MutationObserver", contexts[0].init_scripts[0])
        output.close()

    def test_advisory_mode_writes_suggestion_and_keeps_status_unknown(self):
        source = self.root / "advisory.xlsx"
        rows = [make_row("advisory", claims=[claim("Giao dịch hoàn tất", kind="message")])]
        write_cases(source, rows)
        result = TestcaseRunner(source, browser_factory=lambda _auth: self._fake_browser(self.successful_click),
                                semantic="advisory", judge=lambda *_: {
                                    "verdict": "supported", "evidence_span": "Order saved"}).run()
        output = load_workbook(result["output_path"])
        self.assertEqual(output["Test Cases"]["J6"].value, "Unknown")
        self.assertIn("suggested=Passed", output["Test Cases"]["K6"].value)
        self.assertEqual(result["semantic_suggestions"], 1)
        output.close()

    def test_only_dependent_without_provider_is_prereq_unmet_and_never_runs(self):
        source = self.root / "dependent.xlsx"
        row = make_row("dependent", prerequisites=[{"case_id": "provider"}])
        write_cases(source, [row])
        launched = []
        result = TestcaseRunner(source, only="dependent", semantic="off",
                                browser_factory=lambda _auth: launched.append(1)).run()
        output = load_workbook(result["output_path"])
        self.assertEqual(output["Test Cases"]["J6"].value, "Untested")
        self.assertIn("PREREQ-UNMET:", output["Test Cases"]["K6"].value)
        self.assertIn(output["Test Cases"]["I6"].value, (None, ""))
        self.assertEqual(launched, [])
        output.close()

    def test_fresh_run_does_not_reuse_previous_run_prerequisite_results(self):
        source = self.root / "fresh-run-prereq.xlsx"
        write_cases(source, [make_row("provider"),
                             make_row("dependent", prerequisites=[{"case_id": "provider"}])])
        first = TestcaseRunner(source, browser_factory=lambda _auth: self._fake_browser(self.successful_click),
                               semantic="off").run()
        first_records = [json.loads(line) for line in Path(first["report_path"]).read_text(encoding="utf-8").splitlines()]
        first_run_id = first_records[0]["run_id"]
        launched = []
        second = TestcaseRunner(source, only="dependent", semantic="off",
                                browser_factory=lambda _auth: launched.append(1)).run()
        records = [json.loads(line) for line in Path(second["report_path"]).read_text(encoding="utf-8").splitlines()]
        headers = [record for record in records if record.get("type") == "header"]
        self.assertEqual(len(headers), 2)
        self.assertNotEqual(headers[-1]["run_id"], first_run_id)
        self.assertEqual(launched, [])
        latest_case = next(record for record in reversed(records) if record.get("type") == "case")
        self.assertEqual(latest_case["outcome_code"], "PREREQ-UNMET:")

    @staticmethod
    def _fake_browser(on_click):
        browser = FakeBrowser(on_click)
        context = browser.new_context()
        return browser, context.close, context

    def test_malformed_judge_response_keeps_workbook_and_marks_only_case_unknown(self):
        source = self.root / "malformed.xlsx"
        rows = [make_row("malformed", claims=[claim("Giao dịch xong", kind="other")])]
        write_cases(source, rows)
        browser = FakeBrowser(self.successful_click)
        result = TestcaseRunner(source, browser_factory=lambda _auth: self._fake_browser(self.successful_click),
                                semantic="on", judge=lambda *_: {"bad": "shape"}).run()
        self.assertTrue(Path(result["output_path"]).is_file())
        output = load_workbook(result["output_path"])
        self.assertEqual(output["Test Cases"]["J6"].value, "Unknown")
        output.close()
        records = [json.loads(line) for line in Path(result["report_path"]).read_text(encoding="utf-8").splitlines()]
        case = next(record for record in records if record.get("type") == "case")
        self.assertEqual(case["status"], "Unknown")
        self.assertEqual(case["claim_results"][0]["basis"], "semantic")
        self.assertTrue(case["claim_results"][0]["semantic_error"])

    def test_fixture_secret_is_redacted_from_actual_workbook_and_report(self):
        secret = "token-secret-987654"
        source = self.root / "redacted.xlsx"
        rows = [make_row("redacted", claims=[claim("Saved {{API_TOKEN}}", quoted="Saved {{API_TOKEN}}")])]
        write_cases(source, rows)

        def emit_secret(page):
            page.emit_outcome(f"Saved {secret}")

        browser = FakeBrowser(emit_secret)
        result = TestcaseRunner(source, fixtures={"API_TOKEN": secret},
                                browser_factory=lambda _auth: self._fake_browser(emit_secret), semantic="on",
                                judge=lambda *_: {"malformed": secret}).run()
        output = load_workbook(result["output_path"])
        actual = output["Test Cases"]["I6"].value
        output.close()
        report = Path(result["report_path"]).read_text(encoding="utf-8")
        self.assertNotIn(secret, actual)
        self.assertNotIn(secret, report)

    def test_run_id_tagged_cleanup_runs_after_failed_case_in_separate_context(self):
        source = self.root / "cleanup.xlsx"
        cleanup = [{"tagged_with_run_id": True, "route_actions": [
            {"kind": "click", "role": "button", "name": "Delete test record {{run_id}}", "value": ""}]}]
        rows = [make_row("cleanup", claims=[claim("Order saved", quoted="Order saved")], cleanup=cleanup)]
        write_cases(source, rows)

        def fail_action(page):
            page.emit_outcome("Save failed")

        contexts = []

        def factory(_auth):
            created = self._fake_browser(fail_action)
            contexts.append(created[2])
            return created

        result = TestcaseRunner(source, allow_submit=True, allow_delete=True,
                                browser_factory=factory, semantic="off").run()
        records = [json.loads(line) for line in Path(result["report_path"]).read_text(encoding="utf-8").splitlines()]
        case = next(record for record in records if record.get("type") == "case")
        self.assertEqual(case["status"], "Failed")
        self.assertEqual(case["cleanup"][0]["status"], "completed")
        self.assertTrue(case["cleanup"][0]["observation"])
        self.assertEqual(len(contexts), 2)
        self.assertIsNot(contexts[0], contexts[1])

    def test_calibrate_reports_false_pass_rate_model_and_prompt_version(self):
        from capture_app.workflows.testcase_runner import DEFAULT_CALIBRATION_SET, _run_calibration
        self.assertGreaterEqual(len(DEFAULT_CALIBRATION_SET.read_text(encoding="utf-8").splitlines()), 40)
        self.assertLessEqual(len(DEFAULT_CALIBRATION_SET.read_text(encoding="utf-8").splitlines()), 60)
        labeled = self.root / "calibration.jsonl"
        row = {"claim": claim("Save succeeded", kind="other"),
               "observation": {"messages": ["Save failed"], "action_messages": ["Save failed"],
                               "field_states": [], "headings": [], "url": "https://shop.example.test/",
                               "own_request": {"requests": [], "responses": [], "failures": []}},
               "label": "failed"}
        labeled.write_text(json.dumps(row), encoding="utf-8")
        report = _run_calibration(labeled, judge=lambda *_: {
            "verdict": "supported", "evidence_span": "Save failed"})
        self.assertEqual(report["false_pass_rate"], 1.0)
        self.assertIn("model", report)
        self.assertIn("prompt_version", report)

    def test_results_with_fixture_data_use_fresh_attempt_id(self):
        runner = TestcaseRunner("placeholder.xlsx", semantic="off")
        runner.run_id = "same-run"
        payload = {"start_url": "https://shop.example.test/", "route_actions": [
            {"kind": "fill", "role": "textbox", "name": "Unique key", "value": "{{attempt_id}}"}]}
        first, _ = runner._preflight(payload, "attempt-one")
        second, _ = runner._preflight(payload, "attempt-two")
        self.assertNotEqual(first[0].value, second[0].value)


if __name__ == "__main__":
    unittest.main()
